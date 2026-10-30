"""Sheet-aware XLSX extraction with bounded preflight and Markdown batches."""

from __future__ import annotations

import asyncio
import io
import re
import zipfile
from bisect import bisect_left, bisect_right
from collections.abc import Iterator
from contextlib import closing, contextmanager, nullcontext
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from itertools import chain
from pathlib import Path
from typing import Any

from core.indexing.parsers.registry import parser_registry
from core.indexing.parsers.tabular.base import TabularBatchParser, markdown_row, normalize_headers
from core.indexing.parsers.tabular.xlsx_validation import SheetInfo, TableRegion, XlsxLimits, limit, validate_ooxml
from core.models.document import Document, ProcessedDocument, TextBlock
from core.utils.exceptions import ConfigError, NoIndexableContentError, PipelineError
from core.utils.logging import get_logger

logger = get_logger()


@dataclass(frozen=True)
class _FormulaValue:
    """A cached scalar paired with its inert Excel formula text."""

    cached: Any
    formula: str


_A1_REFERENCE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')*'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(?P<start_col_abs>\$?)(?P<start_col>[A-Z]{1,3})(?P<start_row_abs>\$?)"
    r"(?P<start_row>[1-9][0-9]*)"
    r"(?::(?P<end_col_abs>\$?)(?P<end_col>[A-Z]{1,3})(?P<end_row_abs>\$?)"
    r"(?P<end_row>[1-9][0-9]*))?$",
    re.IGNORECASE,
)
_MAX_FORMULA_SOURCE_COLUMNS = 64


def _rendered_header_key(value: str) -> str:
    """Return the canonical cell text produced by the Markdown renderer."""
    return markdown_row([value])[2:-2]


def _normalize_unique_headers(values: list[str]) -> list[str]:
    """Normalize headers and retain unique Markdown-rendered identities."""
    headers = normalize_headers([value.strip() for value in values])
    reserved = {_rendered_header_key(name) for name in headers}
    used = set()
    for index, name in enumerate(headers):
        candidate, suffix = name, 2
        candidate_key = _rendered_header_key(candidate)
        if candidate_key in used:
            candidate = f"{name}_{suffix}"
            candidate_key = _rendered_header_key(candidate)
            while candidate_key in reserved or candidate_key in used:
                suffix += 1
                candidate = f"{name}_{suffix}"
                candidate_key = _rendered_header_key(candidate)
        headers[index] = candidate
        used.add(candidate_key)
    return headers


def _detect_header(rows, byte_limit, start_row=1):
    """Skip obvious title/preamble rows, retaining ambiguous first rows.

    Look ahead at most 32 physical rows and one batch worth of text. Promote a
    later row only when it is a wider textual row followed by aligned data.
    Equal-width layouts remain ambiguous and keep the first nonempty row.
    Replay buffered rows so physical coordinates remain unchanged.
    """
    buffered = []
    size = 0
    for row in rows:
        buffered.append(row)
        size += sum(len(_cell_text(v).encode("utf-8")) for v in row)
        if len(buffered) >= 32 or size >= byte_limit:
            break
    nonempty = [
        (i, row, [_cell_text(v).strip() for v in row])
        for i, row in enumerate(buffered, start_row)
        if any(_cell_text(v).strip() for v in row)
    ]
    selected = None
    if nonempty:
        selected = nonempty[0][0]
        for index, (number, raw, values) in enumerate(nonempty[1:], 1):
            labels = [value for value in values if value]
            occupied = [column for column, value in enumerate(values) if value]
            earlier = nonempty[:index]
            wider_than_preamble = all(
                sum(bool(value) for value in earlier_values) < len(occupied) for _, _, earlier_values in earlier
            )
            repeated_value_preamble = all(
                len({value for value in earlier_values if value}) == 1 for _, _, earlier_values in earlier
            )
            # Width alone is not enough evidence: a narrow first data record
            # followed by wider records is a valid table and must not lose its
            # first row. A blank physical row is the structural boundary that
            # distinguishes an optional preamble from the table below it.
            separated_from_preamble = number > nonempty[index - 1][0] + 1
            if (
                len(set(labels)) < 2
                or not all(
                    isinstance(value.cached if isinstance(value, _FormulaValue) else value, str)
                    for value in raw
                    if (value.cached if isinstance(value, _FormulaValue) else value) is not None
                )
                or not (wider_than_preamble or repeated_value_preamble)
                or not separated_from_preamble
            ):
                continue
            if index + 1 >= len(nonempty):
                break
            following = nonempty[index + 1][2]
            if all(column < len(following) and following[column] for column in occupied):
                selected = number
                break
    return selected, chain(buffered, rows)


def _column_name(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"


def _column_number(label: str) -> int:
    number = 0
    for char in label.upper():
        number = number * 26 + ord(char) - 64
    return number


def _r1c1_cell(
    column: int,
    row: int,
    *,
    target_column: int,
    target_row: int,
    absolute_column: bool,
    absolute_row: bool,
) -> str:
    if absolute_row:
        rendered_row = f"R{row}"
    else:
        offset = row - target_row
        rendered_row = "R" if offset == 0 else f"R[{offset}]"
    if absolute_column:
        rendered_column = f"C{column}"
    else:
        offset = column - target_column
        rendered_column = "C" if offset == 0 else f"C[{offset}]"
    return rendered_row + rendered_column


def _sheet_name(value: str | None) -> str | None:
    if value is None:
        return None
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1].replace("''", "'")
    return value


def _analyze_formula(
    formula: str,
    *,
    target_row: int,
    target_column: int,
    sheet_name: str,
    visible_columns: list[int],
    headers: list[str],
) -> dict[str, Any]:
    """Normalize inert A1 references and map same-table inputs to headers."""
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import TokenizerError

    header_by_column = dict(zip(visible_columns, headers, strict=True))
    pattern_parts: list[str] = []
    semantic_parts: list[str] = []
    source_columns: list[dict[str, str]] = []
    seen_sources: set[tuple[str, str]] = set()
    all_source_columns: set[int] = set()
    unresolved: list[str] = []
    try:
        tokens = Tokenizer(formula).items
    except TokenizerError:
        tokens = []
    if not tokens:
        return {
            "pattern": formula,
            "semantic_expression": f"{headers[visible_columns.index(target_column)]} = {formula.lstrip('=')}",
            "source_columns": [],
            "unresolved_references": [formula],
        }
    for token in tokens:
        value = token.value
        if token.type != "OPERAND" or token.subtype != "RANGE":
            pattern_parts.append(value)
            semantic_parts.append(value)
            continue
        match = _A1_REFERENCE.fullmatch(value)
        if match is None:
            pattern_parts.append(value)
            semantic_parts.append(value)
            unresolved.append(value)
            continue
        referenced_sheet = _sheet_name(match.group("sheet"))
        prefix = f"{match.group('sheet')}!" if match.group("sheet") else ""
        start_column = _column_number(match.group("start_col"))
        start_row = int(match.group("start_row"))
        start_pattern = _r1c1_cell(
            start_column,
            start_row,
            target_column=target_column,
            target_row=target_row,
            absolute_column=bool(match.group("start_col_abs")),
            absolute_row=bool(match.group("start_row_abs")),
        )
        end_column = start_column
        end_pattern = None
        if match.group("end_col"):
            end_column = _column_number(match.group("end_col"))
            end_pattern = _r1c1_cell(
                end_column,
                int(match.group("end_row")),
                target_column=target_column,
                target_row=target_row,
                absolute_column=bool(match.group("end_col_abs")),
                absolute_row=bool(match.group("end_row_abs")),
            )
        pattern_parts.append(prefix + start_pattern + (f":{end_pattern}" if end_pattern else ""))
        same_sheet = referenced_sheet is None or referenced_sheet == sheet_name
        referenced_headers = []
        if same_sheet:
            lower, upper = sorted((start_column, end_column))
            start_index = bisect_left(visible_columns, lower)
            end_index = bisect_right(visible_columns, upper)
            referenced_columns = visible_columns[start_index:end_index]
            referenced_headers = [header_by_column[column] for column in referenced_columns]
            all_source_columns.update(referenced_columns)
            for column, header in zip(referenced_columns, referenced_headers, strict=True):
                key = (_column_name(column), header)
                if key not in seen_sources and len(source_columns) < _MAX_FORMULA_SOURCE_COLUMNS:
                    source_columns.append({"column": key[0], "header": key[1]})
                    seen_sources.add(key)
        if referenced_headers:
            semantic_parts.append(
                referenced_headers[0]
                if len(referenced_headers) == 1
                else f"{referenced_headers[0]}:{referenced_headers[-1]}"
            )
        else:
            semantic_parts.append(value)
            unresolved.append(value)
    target_header = header_by_column[target_column]
    result = {
        "pattern": "=" + "".join(pattern_parts),
        "semantic_expression": f"{target_header} = {''.join(semantic_parts)}",
        "source_columns": source_columns,
    }
    if len(all_source_columns) > len(source_columns):
        result["source_column_count"] = len(all_source_columns)
        result["source_columns_truncated"] = True
    if unresolved:
        result["unresolved_references"] = list(dict.fromkeys(unresolved))
    return result


def _worksheet_context(sheet_name: str) -> str:
    """Render the worksheet identity as searchable, single-line chunk context."""
    safe_name = sheet_name.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")
    return f"\n\nWorksheet: {safe_name}"


def _cached_cell_text(value: Any) -> str:
    if isinstance(value, _FormulaValue):
        value = value.cached
    if value is None:
        return ""
    elif isinstance(value, bool):
        return "true" if value else "false"
    # OpenPyXL exposes integral numbers and dates with different Python types.
    elif isinstance(value, datetime) and value.time() == time():
        return value.date().isoformat()
    elif isinstance(value, (datetime, date, time)):
        return value.isoformat()
    elif isinstance(value, timedelta):
        return str(value)
    elif isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _cell_text(value: Any) -> str:
    formula = value.formula if isinstance(value, _FormulaValue) else None
    rendered = _cached_cell_text(value)
    if formula is None:
        return rendered
    if rendered:
        return f"{rendered} [formula: {formula}]"
    return f"[formula: {formula}; cached value unavailable]"


def _has_text(values: list[str]) -> bool:
    """Treat cells containing only formatting whitespace as empty."""
    return any(value.strip() for value in values)


@contextmanager
def _document_source(document: Document):
    if document.source_path is not None:
        yield Path(document.source_path)
        return
    if document.raw_bytes is None:
        raise PipelineError("XLSX parsing requires binary workbook content", code="INVALID_XLSX", status_code=422)
    with io.BytesIO(document.raw_bytes) as stream:
        yield stream


def _formula_text(cell) -> str | None:
    """Return stable formula text from regular, shared, or array formulas."""
    if getattr(cell, "data_type", None) != "f":
        return None
    value = cell.value
    text = value if isinstance(value, str) else getattr(value, "text", None)
    if not isinstance(text, str) or not text:
        return None
    return text if text.startswith("=") else f"={text}"


def _rows_with_formulas(value_rows, formula_rows):
    """Merge cached values with formula text while both readers stream."""
    for values, formula_cells in zip(value_rows, formula_rows, strict=True):
        yield tuple(
            _FormulaValue(value, formula) if (formula := _formula_text(cell)) is not None else value
            for value, cell in zip(values, formula_cells, strict=True)
        )


def _single_line(value: Any) -> str:
    return str(value).replace("\r\n", " ").replace("\r", " ").replace("\n", " ")


def _formula_context(groups: list[dict[str, Any]]) -> str:
    if not groups:
        return ""
    lines = ["Formula context (calculated columns):"]
    for group in groups:
        target = _single_line(group["target_header"])
        semantic = _single_line(group["semantic_expression"])
        pattern = _single_line(group["pattern"])
        lines.append(f"- {group['target_column']} ({target}), {group['range']}: {semantic} [pattern: {pattern}]")
    return "\n\n" + "\n".join(lines)


def _compact_formula_batch(
    rows: list[tuple[Any, ...]],
    row_numbers: list[int],
    *,
    sheet_name: str,
    visible_columns: list[int],
    headers: list[str],
    group_repeated: bool = True,
) -> tuple[list[list[str]], list[dict[str, Any]], list[dict[str, Any]], str]:
    """Group repeated relative formulas and retain isolated expressions inline.

    ``group_repeated=False`` is the bounded fallback used when generalized
    context would make an otherwise valid parser batch exceed its byte limit.
    """
    records: list[dict[str, Any]] = []
    for row_number, row in zip(row_numbers, rows, strict=True):
        for position, value in enumerate(row):
            if not isinstance(value, _FormulaValue):
                continue
            target_column_number = visible_columns[position]
            analysis = _analyze_formula(
                value.formula,
                target_row=row_number,
                target_column=target_column_number,
                sheet_name=sheet_name,
                visible_columns=visible_columns,
                headers=headers,
            )
            records.append(
                {
                    "row": row_number,
                    "position": position,
                    "target": f"{_column_name(target_column_number)}{row_number}",
                    "target_column": _column_name(target_column_number),
                    "target_header": headers[position],
                    "formula": value.formula,
                    "cached_value_available": value.cached is not None,
                    **analysis,
                }
            )

    grouped_records: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for record in records:
        sources = tuple((source["column"], source["header"]) for source in record["source_columns"])
        unresolved = tuple(record.get("unresolved_references", []))
        key = (
            record["target_column"],
            record["target_header"],
            record["pattern"],
            record["semantic_expression"],
            sources,
            unresolved,
        )
        grouped_records.setdefault(key, []).append(record)

    candidates: list[tuple[dict[str, Any], list[dict[str, Any]]]] = []
    for matching in grouped_records.values():
        matching.sort(key=lambda record: record["row"])
        segments: list[list[dict[str, Any]]] = []
        for record in matching:
            if not segments or record["row"] != segments[-1][-1]["row"] + 1:
                segments.append([record])
            else:
                segments[-1].append(record)
        for segment in segments:
            if len(segment) < 2:
                continue
            first, last = segment[0], segment[-1]
            group = {
                "target_column": first["target_column"],
                "target_header": first["target_header"],
                "range": f"{first['target_column']}{first['row']}:{first['target_column']}{last['row']}",
                "pattern": first["pattern"],
                "semantic_expression": first["semantic_expression"],
                "source_columns": first["source_columns"],
                "start_row": first["row"],
                "end_row": last["row"],
                "formula_count": len(segment),
            }
            if first.get("unresolved_references"):
                group["unresolved_references"] = first["unresolved_references"]
            if first.get("source_columns_truncated"):
                group["source_column_count"] = first["source_column_count"]
                group["source_columns_truncated"] = True
            candidates.append((group, segment))

    accepted = candidates if group_repeated else []
    formula_groups = [group for group, _ in accepted]
    grouped_targets = {record["target"] for _, segment in accepted for record in segment}
    individual_cells = [
        {key: value for key, value in record.items() if key != "position"}
        for record in records
        if record["target"] not in grouped_targets
    ]
    rendered_rows: list[list[str]] = []
    for row_number, row in zip(row_numbers, rows, strict=True):
        rendered = []
        for position, value in enumerate(row):
            target = f"{_column_name(visible_columns[position])}{row_number}"
            if isinstance(value, _FormulaValue) and target in grouped_targets:
                cached = _cached_cell_text(value)
                rendered.append(cached or "[cached value unavailable]")
            else:
                rendered.append(_cell_text(value))
        rendered_rows.append(rendered)
    context = _formula_context(formula_groups)
    return rendered_rows, formula_groups, individual_cells, context


def _visible_rows(rows, table: TableRegion, info: SheetInfo):
    """Mask hidden cells and expand source-backed vertical row-header merges."""
    visible_columns = [
        column for column in range(table.start_column, table.end_column + 1) if column not in info.hidden_columns
    ]
    visible_offsets = [column - table.start_column for column in visible_columns]
    column_positions = {column: index for index, column in enumerate(visible_columns)}
    merge_starts: dict[int, list[tuple[int, int]]] = {}
    for start_row, end_row, start_column, end_column in info.merged_ranges:
        if (
            start_row < end_row
            and start_column == end_column
            and start_column in column_positions
            and table.start_row <= start_row <= table.end_row
        ):
            merge_starts.setdefault(start_row, []).append((start_column, min(end_row, table.end_row)))
    active_merges: dict[int, tuple[int, Any]] = {}
    width = len(visible_columns)
    for row_number, row in enumerate(rows, start=table.start_row):
        if row_number in info.hidden_rows:
            # Keep a placeholder so downstream physical row coordinates do not shift.
            yield (None,) * width
            continue
        projected = [row[offset] for offset in visible_offsets]
        active_merges = {
            column: (end_row, value) for column, (end_row, value) in active_merges.items() if row_number <= end_row
        }
        for column, (end_row, value) in active_merges.items():
            projected[column_positions[column]] = value
        for column, end_row in merge_starts.get(row_number, []):
            value = projected[column_positions[column]]
            if value is not None:
                active_merges[column] = (end_row, value)
        yield tuple(projected)


def _column_header_hierarchy(
    info: SheetInfo,
    table: TableRegion,
    visible_columns: list[int],
    preamble_values: list[tuple[int, list[str]]],
    leaf_headers: list[str],
) -> tuple[list[str], list[dict[str, Any]], list[int]]:
    """Flatten explicit horizontal merged groups without mistaking a title for one."""
    positions = {column: index for index, column in enumerate(visible_columns)}
    prefixes: list[list[str]] = [[] for _ in visible_columns]
    hierarchy_rows: list[int] = []
    ranges_by_row = {
        row_number: [
            (start_column, end_column)
            for start_row, end_row, start_column, end_column in info.merged_ranges
            if start_row == end_row == row_number
            and start_column < end_column
            and start_column in positions
            and any(column in positions for column in range(start_column, end_column + 1))
        ]
        for row_number, _ in preamble_values
    }
    has_partial_group = any(
        ranges and not (len(ranges) == 1 and ranges[0] == (visible_columns[0], visible_columns[-1]))
        for ranges in ranges_by_row.values()
    )
    for row_number, values in preamble_values:
        ranges = ranges_by_row[row_number]
        # One merged cell spanning the complete table is a title, not a parent
        # for every column unless another row supplies explicit child groups.
        if len(ranges) == 1 and ranges[0] == (visible_columns[0], visible_columns[-1]) and not has_partial_group:
            continue
        applied = False
        for start_column, end_column in ranges:
            label = values[positions[start_column]].strip()
            if not label:
                continue
            for column in visible_columns:
                if start_column <= column <= end_column:
                    prefixes[positions[column]].append(label)
                    applied = True
        if applied:
            hierarchy_rows.append(row_number)
    flattened = []
    hierarchy = []
    for leaf, parents in zip(leaf_headers, prefixes, strict=True):
        path = [*parents, leaf]
        flattened.append(" / ".join(path))
        hierarchy.append({"column": leaf, "path": path})
    return flattened, hierarchy, hierarchy_rows


def _openpyxl_sheets(document: Document, source, infos: list[SheetInfo]):
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    needs_formulas = any(info.formula_cells for info in infos if not info.hidden and info.columns)
    formula_source_manager = _document_source(document) if needs_formulas else nullcontext(None)
    formula_workbook = None
    try:
        with formula_source_manager as formula_source:
            try:
                formula_workbook = (
                    load_workbook(formula_source, read_only=True, data_only=False, keep_links=False)
                    if formula_source is not None
                    else None
                )
                for info in infos:
                    if info.hidden or not info.columns:
                        continue
                    sheet = workbook[info.name]
                    # Exporters may declare A1:A1 or a huge dimension. Use inspected cells.
                    sheet.reset_dimensions()
                    for table in info.tables:
                        with closing(
                            sheet.iter_rows(
                                min_row=table.start_row,
                                max_row=table.end_row,
                                min_col=table.start_column,
                                max_col=table.end_column,
                                values_only=True,
                            )
                        ) as rows:
                            visible_rows = _visible_rows(rows, table, info)
                            if not info.formula_cells:
                                yield info, table, visible_rows
                                continue
                            if formula_workbook is None:
                                raise PipelineError("XLSX formula reader is unavailable", code="INVALID_XLSX")
                            formula_sheet = formula_workbook[info.name]
                            formula_sheet.reset_dimensions()
                            with closing(
                                formula_sheet.iter_rows(
                                    min_row=table.start_row,
                                    max_row=table.end_row,
                                    min_col=table.start_column,
                                    max_col=table.end_column,
                                    values_only=False,
                                )
                            ) as formula_rows:
                                visible_formula_rows = _visible_rows(formula_rows, table, info)
                                yield info, table, _rows_with_formulas(visible_rows, visible_formula_rows)
            finally:
                if formula_workbook is not None:
                    formula_workbook.close()
    finally:
        workbook.close()


@parser_registry.register("xlsx")
class XlsxParser(TabularBatchParser):
    """Extract visible worksheets; limits bound accepted input, not exact RSS."""

    tabular_format = "xlsx"

    def __init__(
        self,
        *,
        batch_size: int = 10_000,
        header_row: int | None = None,
        sheet_header_rows: dict[str, int | None] | None = None,
        sheet_headerless_schemas: dict[str, list[str]] | None = None,
        limits: XlsxLimits | dict[str, int] | None = None,
    ) -> None:
        if type(batch_size) is not int or batch_size < 1:
            raise ConfigError("XLSX batch_size must be a positive integer")
        if header_row is not None and (type(header_row) is not int or not 1 <= header_row <= 1_048_576):
            raise ConfigError("XLSX header_row must be an Excel row number between 1 and 1048576")
        self.header_row = header_row
        if sheet_header_rows is not None and not isinstance(sheet_header_rows, dict):
            raise ConfigError("XLSX sheet_header_rows must map worksheet names to row numbers or null")
        for name, row in (sheet_header_rows or {}).items():
            if not isinstance(name, str) or not name.strip():
                raise ConfigError("XLSX sheet_header_rows requires nonempty worksheet names")
            if row is not None and (type(row) is not int or not 1 <= row <= 1_048_576):
                raise ConfigError("XLSX sheet header rows must be Excel row numbers between 1 and 1048576 or null")
        self.sheet_header_rows = dict(sheet_header_rows or {})
        if sheet_headerless_schemas is not None and not isinstance(sheet_headerless_schemas, dict):
            raise ConfigError("XLSX sheet_headerless_schemas must map worksheet names to column-name lists")
        if header_row is not None and sheet_headerless_schemas:
            raise ConfigError("XLSX header_row cannot be combined with headerless schemas")
        for name, columns in (sheet_headerless_schemas or {}).items():
            if not isinstance(name, str) or not name.strip():
                raise ConfigError("XLSX headerless schemas require nonempty worksheet names")
            if (
                not isinstance(columns, list)
                or not columns
                or len(columns) > 4096
                or any(not isinstance(column, str) or not column.strip() for column in columns)
            ):
                raise ConfigError("XLSX headerless schemas require a nonempty list of nonempty column names")
            base_name = name.rsplit("#", 1)[0]
            if self.sheet_header_rows.get(base_name) is not None:
                raise ConfigError("XLSX sheet header rows cannot be combined with a headerless schema")
        self.sheet_headerless_schemas = {
            name: list(columns) for name, columns in (sheet_headerless_schemas or {}).items()
        }
        self.batch_size = batch_size
        try:
            self.limits = XlsxLimits(**limits) if isinstance(limits, dict) else limits or XlsxLimits()
        except TypeError as exc:
            raise ConfigError(f"Invalid XLSX limits: {exc}") from exc

    def supported_types(self) -> list[str]:
        return ["xlsx"]

    async def parse(self, document: Document) -> ProcessedDocument:
        return await asyncio.to_thread(self._parse, document)

    def iter_batches(self, document: Document) -> Iterator[TextBlock]:
        from openpyxl.utils.exceptions import InvalidFileException

        with _document_source(document) as source:
            infos = validate_ooxml(source, self.limits)
            # Validate overrides before opening a backend or emitting any batches.
            by_name = {info.name: info for info in infos}
            if self.header_row is not None:
                visible = [info for info in infos if not info.hidden and info.columns]
                for info in visible:
                    if self.header_row > info.rows:
                        raise PipelineError(
                            f"XLSX sheet {info.name!r}: selected header row {self.header_row} is missing",
                            code="INVALID_XLSX",
                            status_code=422,
                        )
                    if not any(table.start_row <= self.header_row <= table.end_row for table in info.tables):
                        raise PipelineError(
                            f"XLSX sheet {info.name!r}: selected header row {self.header_row} is empty",
                            code="INVALID_XLSX",
                            status_code=422,
                        )
                    if self.header_row in info.hidden_rows:
                        raise PipelineError(
                            f"XLSX sheet {info.name!r}: selected header row {self.header_row} is hidden",
                            code="INVALID_XLSX",
                            status_code=422,
                        )
            for name, row in self.sheet_header_rows.items():
                info = by_name.get(name)
                if info is None or info.hidden or not info.columns:
                    raise PipelineError(
                        f"XLSX header override targets an absent, hidden, or empty worksheet: {name!r}",
                        code="INVALID_XLSX",
                        status_code=422,
                    )
                if row is not None and row > info.rows:
                    raise PipelineError(
                        f"XLSX sheet {name!r}: selected header row {row} is missing",
                        code="INVALID_XLSX",
                        status_code=422,
                    )
                if row is not None and not any(table.start_row <= row <= table.end_row for table in info.tables):
                    raise PipelineError(
                        f"XLSX sheet {name!r}: selected header row {row} is empty",
                        code="INVALID_XLSX",
                        status_code=422,
                    )
                if row is not None and row in info.hidden_rows:
                    raise PipelineError(
                        f"XLSX sheet {name!r}: selected header row {row} is hidden",
                        code="INVALID_XLSX",
                        status_code=422,
                    )
            with closing(self._iter_sheets(document, source, infos)) as batches:
                try:
                    yield from batches
                except PipelineError:
                    raise
                except (
                    InvalidFileException,
                    OSError,
                    KeyError,
                    IndexError,
                    TypeError,
                    ValueError,
                    SyntaxError,
                    zipfile.BadZipFile,
                ) as exc:
                    raise PipelineError(
                        "File is not a valid XLSX workbook", code="INVALID_XLSX", status_code=422
                    ) from exc

    def _iter_sheets(self, document, source, infos):
        data_rows_seen = 0
        hidden_count = sum(info.hidden for info in infos)
        with closing(_openpyxl_sheets(document, source, infos)) as sheets:
            for sheet, table, row_iterator in sheets:
                visible_columns = [
                    column
                    for column in range(table.start_column, table.end_column + 1)
                    if column not in sheet.hidden_columns
                ]
                visible_column_count = len(visible_columns)
                if visible_column_count == 0:
                    continue
                table_schema_key = f"{sheet.name}#{table.index + 1}"
                headerless_schema = self.sheet_headerless_schemas.get(
                    table_schema_key, self.sheet_headerless_schemas.get(sheet.name)
                )
                configured_header_row = self.sheet_header_rows.get(sheet.name, self.header_row)
                preamble = []
                preamble_rows = []
                preamble_values: list[tuple[int, list[str]]] = []
                preamble_bytes = 0
                header_hierarchy: list[dict[str, Any]] = []
                header_hierarchy_rows: list[int] = []
                header_source = "worksheet"
                if headerless_schema is not None:
                    if len(headerless_schema) != visible_column_count:
                        raise PipelineError(
                            f"XLSX sheet {sheet.name!r} table {table.index + 1}: headerless schema has "
                            f"{len(headerless_schema)} columns, expected {visible_column_count}",
                            code="INVALID_XLSX",
                            status_code=422,
                        )
                    header_values = _normalize_unique_headers(headerless_schema)
                    header_row = None
                    data_row_start = table.start_row
                    header_source = "configured_schema"
                else:
                    selected_header_row = (
                        configured_header_row
                        if configured_header_row is not None
                        and table.start_row <= configured_header_row <= table.end_row
                        else None
                    )
                    if configured_header_row is None and table.header_row_hint is not None:
                        selected_header_row = table.header_row_hint
                    if selected_header_row is None:
                        selected_header_row, row_iterator = _detect_header(
                            row_iterator, self.limits.max_batch_bytes, table.start_row
                        )
                    header_values = None
                    header_row = table.start_row
                    for physical_row, values in enumerate(row_iterator, start=table.start_row):
                        rendered = [_cell_text(value) for value in values]
                        rendered.extend([""] * (visible_column_count - len(rendered)))
                        if selected_header_row is not None and physical_row < selected_header_row:
                            preamble_values.append((physical_row, rendered))
                            values_text = list(dict.fromkeys(value.strip() for value in rendered if value.strip()))
                            if values_text:
                                text = " ; ".join(values_text)
                                preamble_bytes += len(text.encode("utf-8"))
                                if preamble_bytes >= self.limits.max_batch_bytes:
                                    raise limit("XLSX sheet context exceeds the batch byte limit")
                                preamble.append(f"Row {physical_row}: {text}")
                                preamble_rows.append(physical_row)
                            continue
                        if selected_header_row is not None and not _has_text(rendered):
                            raise PipelineError(
                                f"XLSX sheet {sheet.name!r}: selected header row {selected_header_row} is empty",
                                code="INVALID_XLSX",
                                status_code=422,
                            )
                        if _has_text(rendered):
                            if any(len(value) > self.limits.max_cell_chars for value in rendered):
                                raise limit("XLSX header text exceeds the character limit")
                            if sum(len(value.encode("utf-8")) for value in rendered) >= self.limits.max_batch_bytes:
                                raise limit("XLSX header exceeds the batch byte limit")
                            leaf_headers = normalize_headers([value.strip() for value in rendered])
                            flattened, header_hierarchy, header_hierarchy_rows = _column_header_hierarchy(
                                sheet, table, visible_columns, preamble_values, leaf_headers
                            )
                            header_values = _normalize_unique_headers(flattened)
                            for index, hierarchy in enumerate(header_hierarchy):
                                hierarchy["column"] = header_values[index]
                            header_row = physical_row
                            break
                    if not header_values:
                        if selected_header_row is not None:
                            raise PipelineError(
                                f"XLSX sheet {sheet.name!r}: selected header row {selected_header_row} is missing",
                                code="INVALID_XLSX",
                                status_code=422,
                            )
                        continue
                    data_row_start = header_row + 1
                if sheet.missing_formula_values:
                    logger.bind(sheet_name=sheet.name, missing_formula_values=sheet.missing_formula_values).warning(
                        "XLSX formula text was indexed without a cached result; formulas are not recalculated"
                    )
                header_lines = [markdown_row(header_values), markdown_row(["---"] * len(header_values))]
                worksheet_context = _worksheet_context(sheet.name)
                sheet_context = (
                    "\n\nWorksheet context (before column headers):\n" + "\n".join(preamble) if preamble else ""
                )
                context = worksheet_context + sheet_context
                header_bytes = len(("\n".join(header_lines) + context).encode("utf-8"))
                if header_bytes >= self.limits.max_batch_bytes:
                    raise limit("XLSX header exceeds the batch byte limit")
                batch_rows: list[tuple[Any, ...]] = []
                row_numbers, batch_bytes = [], header_bytes
                row_header_positions = sorted(
                    {
                        visible_columns.index(start_column)
                        for start_row, end_row, start_column, end_column in sheet.merged_ranges
                        if start_row < end_row
                        and start_column == end_column
                        and start_column in visible_columns
                        and start_row >= data_row_start
                        and start_row <= table.end_row
                    }
                )
                row_header_paths: dict[str, list[str]] = {}

                def block():
                    def render(group_repeated: bool):
                        rendered_rows, formula_groups, formula_cells, formula_context = _compact_formula_batch(
                            batch_rows,
                            row_numbers,
                            sheet_name=sheet.name,
                            visible_columns=visible_columns,
                            headers=header_values,
                            group_repeated=group_repeated,
                        )
                        lines = [*header_lines, *(markdown_row(row) for row in rendered_rows)]
                        block_bytes = len(
                            ("\n".join(lines) + worksheet_context + sheet_context + formula_context).encode("utf-8")
                        )
                        return lines, formula_groups, formula_cells, formula_context, block_bytes

                    lines, formula_groups, formula_cells, formula_context, block_bytes = render(True)
                    if block_bytes > self.limits.max_batch_bytes:
                        # Generalized formula context is normally smaller than
                        # repeated inline expressions, but short runs can add
                        # overhead. Preserve ingestion and formula information
                        # by reverting this bounded batch to individual cells.
                        lines, formula_groups, formula_cells, formula_context, block_bytes = render(False)
                    if block_bytes > self.limits.max_batch_bytes:
                        raise limit("XLSX batch exceeds the batch byte limit")
                    return self._block(
                        document,
                        sheet,
                        table,
                        hidden_count,
                        header_values,
                        header_row,
                        lines,
                        row_numbers,
                        worksheet_context,
                        sheet_context,
                        preamble_rows,
                        header_source,
                        header_hierarchy,
                        header_hierarchy_rows,
                        [header_values[index] for index in row_header_positions],
                        row_header_paths,
                        formula_groups,
                        formula_cells,
                        formula_context,
                    )

                for physical_row, values in enumerate(row_iterator, start=data_row_start):
                    rendered = [_cell_text(value) for value in values]
                    if not _has_text(rendered):
                        continue
                    if any(len(value) > self.limits.max_cell_chars for value in rendered):
                        raise limit("XLSX cell text exceeds the character limit")
                    rendered.extend([""] * (len(header_values) - len(rendered)))
                    if len(rendered) > len(header_values):
                        raise PipelineError("XLSX reader returned more columns than inspected", code="INVALID_XLSX")
                    if (
                        header_bytes + sum(len(value.encode("utf-8")) for value in rendered)
                        >= self.limits.max_batch_bytes
                    ):
                        raise limit(f"XLSX row {physical_row} exceeds the batch byte limit")
                    row = markdown_row(rendered)
                    path = [rendered[index] for index in row_header_positions if rendered[index]]
                    row_bytes = len(row.encode("utf-8")) + 1 + sum(len(value.encode("utf-8")) for value in path)
                    if header_bytes + row_bytes > self.limits.max_batch_bytes:
                        raise limit(f"XLSX row {physical_row} exceeds the batch byte limit")
                    if row_numbers and batch_bytes + row_bytes > self.limits.max_batch_bytes:
                        yield block()
                        batch_rows, row_numbers, batch_bytes = [], [], header_bytes
                        row_header_paths = {}
                    batch_rows.append(tuple(values))
                    row_numbers.append(physical_row)
                    if path:
                        row_header_paths[str(physical_row)] = path
                    batch_bytes += row_bytes
                    data_rows_seen += 1
                    if len(row_numbers) == self.batch_size:
                        yield block()
                        batch_rows, row_numbers, batch_bytes = [], [], header_bytes
                        row_header_paths = {}
                if row_numbers:
                    yield block()
        if not data_rows_seen:
            raise NoIndexableContentError("XLSX contains no visible worksheet with indexable data rows")

    def _block(
        self,
        document,
        sheet,
        table: TableRegion,
        hidden_count,
        headers,
        header_row,
        lines,
        row_numbers,
        worksheet_context,
        sheet_context,
        preamble_rows,
        header_source,
        header_hierarchy,
        header_hierarchy_rows,
        row_header_columns,
        row_header_paths,
        formula_groups,
        formula_cells,
        formula_context,
    ):
        sheet_ref = "'" + sheet.name.replace("'", "''") + "'!"
        visible_columns = [
            column for column in range(table.start_column, table.end_column + 1) if column not in sheet.hidden_columns
        ]
        first_column = _column_name(visible_columns[0])
        last_column = _column_name(visible_columns[-1])
        metadata = {
            "tabular_format": "xlsx",
            "workbook_filename": document.filename,
            "sheet_name": sheet.name,
            "worksheet_context": worksheet_context,
            "sheet_index": sheet.index,
            "table_id": f"sheet:{sheet.index}:table:{table.index}",
            "table_index": table.index,
            "worksheet_table_count": len(sheet.tables),
            "table_columns": list(headers),
            "table_header_row": header_row,
            "table_header_source": header_source,
            "table_start_row": header_row if header_row is not None else table.start_row,
            "table_end_row": table.end_row,
            "table_start_column": visible_columns[0],
            "table_end_column": visible_columns[-1],
            "table_worksheet_column_indices": visible_columns,
            "table_range": (
                f"{sheet_ref}{first_column}{header_row}:{last_column}{table.end_row}"
                if header_row is not None
                else f"{sheet_ref}{first_column}{table.start_row}:{last_column}{table.end_row}"
            ),
            "table_region_range": f"{sheet_ref}{first_column}{table.start_row}:{last_column}{table.end_row}",
            "table_region_start_row": table.start_row,
            "table_row_start": row_numbers[0],
            "table_row_end": row_numbers[-1],
            "table_row_numbers": list(row_numbers),
            "cell_range": f"{sheet_ref}{first_column}{row_numbers[0]}:{last_column}{row_numbers[-1]}",
            "table_header_range": (
                f"{sheet_ref}{first_column}{header_row}:{last_column}{header_row}" if header_row is not None else None
            ),
            "hidden_sheet": False,
            "hidden_sheets_skipped": hidden_count,
            "hidden_rows_skipped": len(sheet.hidden_rows),
            "hidden_columns_skipped": len(sheet.hidden_columns),
            "formula_policy": "compact_groups_and_individual_formulas",
            "sheet_formula_cells": sheet.formula_cells,
            "sheet_missing_formula_values": sheet.missing_formula_values,
            "spreadsheet_backend": "openpyxl",
            # Transitional aliases keep existing chunking/retrieval operational.
            "csv_columns": list(headers),
            "csv_row_start": row_numbers[0],
            "csv_row_end": row_numbers[-1],
            "csv_row_numbers": list(row_numbers),
            "csv_bad_rows": 0,
        }
        if header_hierarchy_rows:
            metadata.update(
                table_header_hierarchy=header_hierarchy,
                table_header_hierarchy_rows=list(header_hierarchy_rows),
            )
        if row_header_columns:
            metadata.update(
                table_row_header_columns=list(row_header_columns),
                table_row_header_paths=dict(row_header_paths),
                table_merged_row_headers_expanded=True,
            )
        if formula_groups:
            metadata["formula_groups"] = list(formula_groups)
        if formula_cells:
            metadata["formula_cells"] = list(formula_cells)
        if formula_context:
            metadata["formula_context"] = formula_context
        if sheet_context:
            metadata.update(sheet_context=sheet_context, sheet_context_row_numbers=list(preamble_rows))
        return TextBlock(
            text="\n".join(lines) + worksheet_context + sheet_context + formula_context,
            block_type="table",
            metadata=metadata,
        )

    def _parse(self, document: Document) -> ProcessedDocument:
        # The compatibility API collects blocks; production ingestion uses iter_batches.
        return ProcessedDocument(
            document_id=document.id, text_blocks=list(self.iter_batches(document)), metadata=dict(document.metadata)
        )
