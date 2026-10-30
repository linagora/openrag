from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from core.indexing.parsers.tabular import xlsx_parser as module
from core.indexing.parsers.tabular.xlsx_parser import XlsxLimits, XlsxParser
from core.indexing.parsers.tabular.xlsx_validation import MAIN_TYPE
from core.models.document import Document, DocumentType
from core.utils.exceptions import NoIndexableContentError, PipelineError
from openpyxl import Workbook


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    people = workbook.active
    people.title = "People"
    people.append([])
    people.append(["name", "", "active", "joined"])
    people.append(["Alice", "Paris", True, date(2025, 1, 2)])
    people.append([])
    people.append(["Bob", "Lyon", False, date(2025, 2, 3)])

    projects = workbook.create_sheet("Projects")
    projects.append(["name", "summary", ""])
    projects.append(["Access", "Public-service accessibility", "owner@example.org"])

    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["secret"])
    hidden.append(["must not be indexed"])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_extracts_visible_sheets_with_physical_rows_and_stable_metadata():
    parser = XlsxParser(batch_size=1)
    document = Document(filename="people.xlsx", content_type=DocumentType.XLSX, raw_bytes=_workbook_bytes())

    blocks = list(parser.iter_batches(document))

    assert [block.metadata["sheet_name"] for block in blocks] == ["People", "People", "Projects"]
    assert [block.metadata["table_row_numbers"] for block in blocks] == [[3], [5], [2]]
    assert blocks[0].metadata["table_columns"] == ["name", "column_2", "active", "joined"]
    assert blocks[0].metadata["table_id"] == "sheet:0:table:0"
    assert blocks[2].metadata["table_id"] == "sheet:1:table:0"
    assert blocks[2].metadata["table_columns"] == ["name", "summary", "column_3"]
    assert blocks[0].metadata["cell_range"] == "'People'!A3:D3"
    assert blocks[0].metadata["worksheet_context"] == "\n\nWorksheet: People"
    assert blocks[2].metadata["worksheet_context"] == "\n\nWorksheet: Projects"
    assert "Worksheet: People" in blocks[0].text
    assert "Worksheet: Projects" in blocks[2].text
    assert "2025-01-02" in blocks[0].text
    assert "true" in blocks[0].text
    assert "must not be indexed" not in "\n".join(block.text for block in blocks)


def test_reader_is_openpyxl_only():
    [block] = list(XlsxParser().iter_batches(_document([["name"], ["Alice"]])))

    assert block.metadata["spreadsheet_backend"] == "openpyxl"
    assert "spreadsheet_fallback_used" not in block.metadata


def test_worksheet_name_is_searchable_without_a_table_preamble():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Research O'Brien"
    sheet.append(["dataset_id", "dataset"])
    sheet.append(["DAT-001", "Public services"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))
    [chunk] = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata),
        chunk_size=100,
        length_function=lambda text: len(text.split()),
    )

    expected_context = "\n\nWorksheet: Research O'Brien"
    assert block.metadata["sheet_name"] == "Research O'Brien"
    assert block.metadata["worksheet_context"] == expected_context
    assert block.text.endswith(expected_context)
    assert chunk.content.endswith(expected_context)
    assert chunk.metadata["cell_range"] == "'Research O''Brien'!A2:B2"


def test_detects_and_extracts_each_table_region_with_exact_coordinates():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"
    for row in (("person_id", "name"), ("P-1", "Alice")):
        sheet.append(row)
    sheet["D1"], sheet["E1"] = "office_id", "city"
    sheet["D2"], sheet["E2"] = "O-1", "Paris"
    sheet["L30"], sheet["M30"] = "metric", "value"
    sheet["L31"], sheet["M31"] = "latency", 42
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert [block.metadata["table_id"] for block in blocks] == [
        "sheet:0:table:0",
        "sheet:0:table:1",
        "sheet:0:table:2",
    ]
    assert all(block.metadata["worksheet_table_count"] == 3 for block in blocks)
    assert [block.metadata["table_range"] for block in blocks] == [
        "'Dashboard'!A1:B2",
        "'Dashboard'!D1:E2",
        "'Dashboard'!L30:M31",
    ]
    assert [block.metadata["table_header_range"] for block in blocks] == [
        "'Dashboard'!A1:B1",
        "'Dashboard'!D1:E1",
        "'Dashboard'!L30:M30",
    ]
    assert [block.metadata["table_row_numbers"] for block in blocks] == [[2], [2], [31]]
    assert "latency" in blocks[2].text
    chunks = [
        chunk
        for block in blocks
        for chunk in chunk_table(
            MDElement(type="table", content=block.text, metadata=block.metadata),
            chunk_size=100,
            length_function=lambda text: len(text.split()),
        )
    ]
    assert [chunk.metadata["table_id"] for chunk in chunks] == [
        "sheet:0:table:0",
        "sheet:0:table:1",
        "sheet:0:table:2",
    ]


def test_chunked_offset_table_keeps_absolute_worksheet_citations():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet["L30"], sheet["M30"] = "metric", "value"
    sheet["L31"], sheet["M31"] = "latency", "long value " * 20
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    chunks = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata),
        chunk_size=20,
        length_function=lambda text: len(text.split()),
    )

    assert len(chunks) > 1
    ranges = {chunk.metadata["cell_range"] for chunk in chunks}
    assert ranges == {"'Sheet'!L31:L31", "'Sheet'!M31:M31"}
    assert all(chunk.metadata["table_id"] == "sheet:0:table:0" for chunk in chunks)


def test_merged_table_titles_are_context_not_column_headers():
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "People"
    sheet.append(["id", "name"])
    sheet.append(["P-1", "Alice"])
    sheet.append([])
    sheet.merge_cells("A5:B5")
    sheet["A5"] = "Offices"
    sheet.append(["id", "city"])
    sheet.append(["O-1", "Paris"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert [block.metadata["table_columns"] for block in blocks] == [["id", "name"], ["id", "city"]]
    assert [block.metadata["table_header_row"] for block in blocks] == [2, 6]
    assert [block.metadata["table_range"] for block in blocks] == ["'Sheet'!A2:B3", "'Sheet'!A6:B7"]
    assert [block.metadata["table_region_range"] for block in blocks] == ["'Sheet'!A1:B3", "'Sheet'!A5:B7"]
    assert [block.metadata["sheet_context_row_numbers"] for block in blocks] == [[1], [5]]
    assert [block.metadata["table_row_numbers"] for block in blocks] == [[3], [7]]


def test_merged_column_header_groups_are_flattened_with_hierarchy_metadata():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Directory"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Person"
    sheet.merge_cells("C1:D1")
    sheet["C1"] = "Contact"
    sheet.append(["id", "name", "allowed", "email"])
    sheet.append(["P-1", "Alice", True, "alice@example.org"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_columns"] == [
        "Person / id",
        "Person / name",
        "Contact / allowed",
        "Contact / email",
    ]
    assert block.metadata["table_header_hierarchy_rows"] == [1]
    assert block.metadata["table_header_hierarchy"] == [
        {"column": "Person / id", "path": ["Person", "id"]},
        {"column": "Person / name", "path": ["Person", "name"]},
        {"column": "Contact / allowed", "path": ["Contact", "allowed"]},
        {"column": "Contact / email", "path": ["Contact", "email"]},
    ]
    assert block.metadata["table_header_row"] == 2
    assert block.metadata["table_row_numbers"] == [3]
    assert "| Person / id | Person / name | Contact / allowed | Contact / email |" in block.text


def test_nested_merged_column_headers_select_leaf_row_and_preserve_full_paths():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Directory"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Directory"
    sheet.merge_cells("A2:B2")
    sheet["A2"] = "Person"
    sheet.merge_cells("C2:D2")
    sheet["C2"] = "Contact"
    sheet.append(["id", "name", "allowed", "email"])
    sheet.append(["P-1", "Alice", True, "alice@example.org"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_header_row"] == 3
    assert block.metadata["table_header_hierarchy_rows"] == [1, 2]
    assert block.metadata["table_columns"] == [
        "Directory / Person / id",
        "Directory / Person / name",
        "Directory / Contact / allowed",
        "Directory / Contact / email",
    ]
    assert block.metadata["table_header_hierarchy"] == [
        {"column": "Directory / Person / id", "path": ["Directory", "Person", "id"]},
        {"column": "Directory / Person / name", "path": ["Directory", "Person", "name"]},
        {
            "column": "Directory / Contact / allowed",
            "path": ["Directory", "Contact", "allowed"],
        },
        {
            "column": "Directory / Contact / email",
            "path": ["Directory", "Contact", "email"],
        },
    ]
    assert block.metadata["table_row_numbers"] == [4]


def test_merged_header_hierarchy_is_scoped_to_each_separated_table():
    workbook = Workbook()
    sheet = workbook.active
    for top_row, prefix, record in ((1, "Person", "P-1"), (6, "Office", "O-1")):
        sheet.merge_cells(start_row=top_row, start_column=1, end_row=top_row, end_column=2)
        sheet.cell(top_row, 1, prefix)
        sheet.merge_cells(start_row=top_row, start_column=3, end_row=top_row, end_column=4)
        sheet.cell(top_row, 3, "Details")
        sheet.cell(top_row + 1, 1, "id")
        sheet.cell(top_row + 1, 2, "name")
        sheet.cell(top_row + 1, 3, "status")
        sheet.cell(top_row + 1, 4, "summary")
        sheet.cell(top_row + 2, 1, record)
        sheet.cell(top_row + 2, 2, prefix)
        sheet.cell(top_row + 2, 3, "active")
        sheet.cell(top_row + 2, 4, f"{prefix} summary")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 2
    assert blocks[0].metadata["table_columns"][0] == "Person / id"
    assert blocks[1].metadata["table_columns"][0] == "Office / id"
    assert blocks[0].metadata["table_id"] == "sheet:0:table:0"
    assert blocks[1].metadata["table_id"] == "sheet:0:table:1"
    assert "Office / id" not in blocks[0].text
    assert "Person / id" not in blocks[1].text


def test_vertical_merged_row_headers_are_expanded_and_paths_follow_chunks():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["region", "country", "amount"])
    sheet["A2"] = "Europe"
    sheet["B2"] = "France"
    sheet["C2"] = 10
    sheet["B3"] = "Germany"
    sheet["C3"] = 20
    sheet.merge_cells("A2:A3")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser(batch_size=1).iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 2
    assert all("Europe" in block.text for block in blocks)
    assert all(block.metadata["table_row_header_columns"] == ["region"] for block in blocks)
    assert blocks[0].metadata["table_row_header_paths"] == {"2": ["Europe"]}
    assert blocks[1].metadata["table_row_header_paths"] == {"3": ["Europe"]}
    chunks = [
        chunk
        for block in blocks
        for chunk in chunk_table(
            MDElement(type="table", content=block.text, metadata=block.metadata),
            chunk_size=100,
            length_function=lambda text: len(text.split()),
        )
    ]
    assert [chunk.metadata["table_row_header_paths"] for chunk in chunks] == [
        {"2": ["Europe"]},
        {"3": ["Europe"]},
    ]


def test_nested_vertical_merged_row_headers_preserve_hierarchical_paths():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["region", "country", "city", "amount"])
    rows = [
        ["Europe", "France", "Paris", 10],
        [None, None, "Lyon", 20],
        [None, "Germany", "Berlin", 30],
        [None, None, "Hamburg", 40],
    ]
    for row in rows:
        sheet.append(row)
    sheet.merge_cells("A2:A5")
    sheet.merge_cells("B2:B3")
    sheet.merge_cells("B4:B5")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_row_header_columns"] == ["region", "country"]
    assert block.metadata["table_row_header_paths"] == {
        "2": ["Europe", "France"],
        "3": ["Europe", "France"],
        "4": ["Europe", "Germany"],
        "5": ["Europe", "Germany"],
    }
    assert "| Europe | France | Paris | 10 |" in block.text
    assert "| Europe | France | Lyon | 20 |" in block.text
    assert "| Europe | Germany | Berlin | 30 |" in block.text
    assert "| Europe | Germany | Hamburg | 40 |" in block.text


def test_merged_row_header_expansion_is_scoped_to_separated_tables():
    workbook = Workbook()
    sheet = workbook.active
    for header_row, region, amounts in ((1, "Europe", (10, 20)), (5, "Asia", (30, 40))):
        sheet.cell(header_row, 1, "region")
        sheet.cell(header_row, 2, "amount")
        sheet.cell(header_row + 1, 1, region)
        sheet.cell(header_row + 1, 2, amounts[0])
        sheet.cell(header_row + 2, 2, amounts[1])
        sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 2,
            end_column=1,
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 2
    assert blocks[0].metadata["table_row_header_paths"] == {"2": ["Europe"], "3": ["Europe"]}
    assert blocks[1].metadata["table_row_header_paths"] == {"6": ["Asia"], "7": ["Asia"]}
    assert "Asia" not in blocks[0].text
    assert "Europe" not in blocks[1].text


def test_headerless_schema_preserves_first_row_as_data():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Raw"
    sheet.append(["HL-001", "Alice", "Accessibility"])
    sheet.append(["HL-002", "René", "Search"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(
        XlsxParser(sheet_headerless_schemas={"Raw": ["id", "name", "topic"]}).iter_batches(
            Document(raw_bytes=output.getvalue())
        )
    )

    assert block.metadata["table_columns"] == ["id", "name", "topic"]
    assert block.metadata["table_header_row"] is None
    assert block.metadata["table_header_range"] is None
    assert block.metadata["table_header_source"] == "configured_schema"
    assert block.metadata["table_row_numbers"] == [1, 2]
    assert "HL-001" in block.text


def test_table_specific_headerless_schemas_work_across_tables_and_sheets():
    workbook = Workbook()
    first = workbook.active
    first.title = "Raw"
    first.append(["P-1", "Alice"])
    first.append([])
    first.append(["O-1", "Paris", "Support"])
    second = workbook.create_sheet("Metrics")
    second.append(["M-1", 42])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(
        XlsxParser(
            sheet_headerless_schemas={
                "Raw#1": ["person_id", "name"],
                "Raw#2": ["office_id", "city", "service"],
                "Metrics": ["metric_id", "value"],
            }
        ).iter_batches(Document(raw_bytes=output.getvalue()))
    )

    assert [block.metadata["table_columns"] for block in blocks] == [
        ["person_id", "name"],
        ["office_id", "city", "service"],
        ["metric_id", "value"],
    ]
    assert [block.metadata["sheet_name"] for block in blocks] == ["Raw", "Raw", "Metrics"]
    assert [block.metadata["table_id"] for block in blocks] == [
        "sheet:0:table:0",
        "sheet:0:table:1",
        "sheet:1:table:0",
    ]


def test_detects_extracts_and_chunks_ten_tables_in_one_worksheet():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ten Tables"
    expected_ranges = []
    for table_index in range(10):
        header_row = table_index * 3 + 1
        data_row = header_row + 1
        sheet.cell(header_row, 1, "id")
        sheet.cell(header_row, 2, "value")
        sheet.cell(data_row, 1, f"T-{table_index + 1:02d}")
        sheet.cell(data_row, 2, table_index + 1)
        expected_ranges.append(f"'Ten Tables'!A{header_row}:B{data_row}")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 10
    assert [block.metadata["table_index"] for block in blocks] == list(range(10))
    assert [block.metadata["table_range"] for block in blocks] == expected_ranges
    assert all(block.metadata["worksheet_table_count"] == 10 for block in blocks)
    assert all(block.metadata["sheet_name"] == "Ten Tables" for block in blocks)
    chunks = [
        chunk
        for block in blocks
        for chunk in chunk_table(
            MDElement(type="table", content=block.text, metadata=block.metadata),
            chunk_size=100,
            length_function=lambda text: len(text.split()),
        )
    ]
    assert len(chunks) == 10
    assert [chunk.metadata["table_id"] for chunk in chunks] == [f"sheet:0:table:{index}" for index in range(10)]
    assert [chunk.metadata["cell_range"] for chunk in chunks] == [
        f"'Ten Tables'!A{index * 3 + 2}:B{index * 3 + 2}" for index in range(10)
    ]


def test_multiple_tables_across_sheets_keep_sheet_identity_and_context_when_chunked():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    workbook.active.title = "North Region"
    workbook.create_sheet("South Region")
    expected = []
    for sheet_index, sheet in enumerate(workbook.worksheets):
        for table_index in range(2):
            title_row = table_index * 4 + 1
            header_row = title_row + 1
            data_row = title_row + 2
            title = f"{sheet.title} table {table_index + 1} context"
            sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=2)
            sheet.cell(title_row, 1, title)
            sheet.cell(header_row, 1, "id")
            sheet.cell(header_row, 2, "value")
            sheet.cell(data_row, 1, f"S{sheet_index}-T{table_index}")
            sheet.cell(data_row, 2, title)
            expected.append((sheet_index, sheet.title, table_index, title, data_row))
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 4
    for block, (sheet_index, sheet_name, table_index, title, data_row) in zip(blocks, expected, strict=True):
        assert block.metadata["sheet_index"] == sheet_index
        assert block.metadata["sheet_name"] == sheet_name
        assert block.metadata["table_index"] == table_index
        assert block.metadata["table_id"] == f"sheet:{sheet_index}:table:{table_index}"
        assert block.metadata["worksheet_table_count"] == 2
        assert block.metadata["worksheet_context"] == f"\n\nWorksheet: {sheet_name}"
        assert block.metadata["sheet_context_row_numbers"] == [data_row - 2]
        assert title in block.metadata["sheet_context"]
        assert f"Worksheet: {sheet_name}" in block.text
        assert title in block.text

    chunks = [
        chunk
        for block in blocks
        for chunk in chunk_table(
            MDElement(type="table", content=block.text, metadata=block.metadata),
            chunk_size=100,
            length_function=lambda text: len(text.split()),
        )
    ]
    assert len(chunks) == 4
    for chunk, (_, sheet_name, table_index, title, _) in zip(chunks, expected, strict=True):
        assert chunk.metadata["sheet_name"] == sheet_name
        assert chunk.metadata["table_index"] == table_index
        assert chunk.metadata["worksheet_context"] == f"\n\nWorksheet: {sheet_name}"
        assert title in chunk.metadata["sheet_context"]
        assert f"Worksheet: {sheet_name}" in chunk.content
        assert title in chunk.content


def test_blank_cells_inside_a_table_do_not_split_its_region():
    document = _document(
        [
            ["id", "optional", "summary"],
            ["A-1", None, "First"],
            ["A-2", "present", "Second"],
        ]
    )

    [block] = list(XlsxParser().iter_batches(document))

    assert block.metadata["worksheet_table_count"] == 1
    assert block.metadata["table_range"] == "'Sheet'!A1:C3"
    assert block.metadata["table_row_numbers"] == [2, 3]


def test_style_only_cells_do_not_create_false_table_regions():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "value"])
    sheet.append([1, "kept"])
    sheet["Z100"].number_format = "0.00"
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["worksheet_table_count"] == 1
    assert block.metadata["table_range"] == "'Sheet'!A1:B2"


def test_rejects_excessive_detected_table_fragmentation_before_reader_opens(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    for column in (1, 4, 7):
        sheet.cell(1, column, "id")
        sheet.cell(2, column, column)
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))
    with pytest.raises(PipelineError, match="table count exceeds") as error:
        list(XlsxParser(limits={"max_tables": 2}).iter_batches(Document(raw_bytes=output.getvalue())))
    assert error.value.code == "XLSX_LIMIT_EXCEEDED"


def test_rejects_non_xlsx_archives_and_resource_amplification():
    document = Document(filename="fake.xlsx", content_type=DocumentType.XLSX, raw_bytes=b"not a zip")
    with pytest.raises(PipelineError, match="valid XLSX"):
        list(XlsxParser().iter_batches(document))

    document = Document(filename="people.xlsx", content_type=DocumentType.XLSX, raw_bytes=_workbook_bytes())
    parser = XlsxParser(limits=XlsxLimits(max_uncompressed_bytes=1))
    with pytest.raises(PipelineError, match="extraction limit"):
        list(parser.iter_batches(document))


def test_rejects_workbook_without_visible_data_rows():
    workbook = Workbook()
    workbook.active.append(["header"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    document = Document(filename="empty.xlsx", content_type=DocumentType.XLSX, raw_bytes=output.getvalue())
    with pytest.raises(NoIndexableContentError, match="no visible worksheet"):
        list(XlsxParser().iter_batches(document))


def test_names_a_trailing_header_cell_when_its_column_contains_data():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name"])
    sheet.append(["Alice", "unexpected"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    document = Document(filename="wide.xlsx", content_type=DocumentType.XLSX, raw_bytes=output.getvalue())
    [block] = list(XlsxParser().iter_batches(document))

    assert block.metadata["table_columns"] == ["name", "column_2"]
    assert "unexpected" in block.text


def _document(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return Document(filename="sample.xlsx", content_type=DocumentType.XLSX, raw_bytes=output.getvalue())


def _edit_xml(document, edit, target="xl/worksheets/sheet1.xml"):
    import zipfile

    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(document.raw_bytes)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            data = source.read(name)
            if name == target:
                data = edit(data.decode()).encode()
            dest.writestr(name, data)
    return document.model_copy(update={"raw_bytes": output.getvalue()})


def _replace_part(document, target, data):
    import zipfile

    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(document.raw_bytes)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            dest.writestr(name, data if name == target else source.read(name))
    return document.model_copy(update={"raw_bytes": output.getvalue()})


def _transcode_xml_parts(document, xml_encoding):
    """Rewrite OOXML text parts with an explicit encoding and BOM."""
    import codecs
    import zipfile

    codec, declaration_encoding, bom = {
        "utf-8-bom": ("utf-8", "UTF-8", codecs.BOM_UTF8),
        "utf-16-le": ("utf-16-le", "UTF-16", codecs.BOM_UTF16_LE),
        "utf-16-be": ("utf-16-be", "UTF-16", codecs.BOM_UTF16_BE),
    }[xml_encoding]
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(document.raw_bytes)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            data = source.read(name)
            if name.endswith((".xml", ".rels")):
                text = data.decode("utf-8")
                if text.startswith("<?xml"):
                    text = text.split("?>", 1)[1]
                declaration = f'<?xml version="1.0" encoding="{declaration_encoding}"?>'
                data = bom + (declaration + text).encode(codec)
            dest.writestr(name, data)
    return document.model_copy(update={"raw_bytes": output.getvalue()})


def test_accepts_workbook_content_type_declared_as_default():
    from xml.etree import ElementTree as ET

    def use_default(xml):
        root = ET.fromstring(xml)
        for node in list(root):
            if node.get("PartName") == "/xl/workbook.xml":
                root.remove(node)
            elif node.get("Extension", "").lower() == "xml":
                node.set("ContentType", MAIN_TYPE)
        return ET.tostring(root, encoding="unicode")

    doc = _edit_xml(_document([["name"], ["Alice"]]), use_default, target="[Content_Types].xml")

    [block] = list(XlsxParser().iter_batches(doc))

    assert block.metadata["table_row_numbers"] == [2]
    assert "Alice" in block.text


def test_renders_numbers_dates_unicode_and_blank_columns():
    from datetime import datetime, time

    doc = _document(
        [
            [None, "id", "date", "timestamp", "time", "active", "text"],
            [None, 12, date(2025, 1, 2), datetime(2025, 1, 2, 12, 30), time(9, 15), True, "café René 東京"],
        ]
    )
    [block] = list(XlsxParser().iter_batches(doc))
    assert "| 12 | 2025-01-02 | 2025-01-02T12:30:00 | 09:15:00 | true | café René 東京 |" in block.text
    assert block.metadata["table_start_column"] == 2
    assert block.metadata["table_header_range"] == "'Sheet'!B1:G1"


@pytest.mark.parametrize("xml_encoding", ["utf-8-bom", "utf-16-le", "utf-16-be"])
@pytest.mark.parametrize("use_shared_strings", [False, True])
def test_preserves_unicode_from_bom_marked_ooxml_parts(xml_encoding, use_shared_strings):
    rows = [
        ["kind", "value"],
        ["composed", "café"],
        ["decomposed", "cafe\u0301"],
        ["right-to-left", "مرحبا שלום"],
        ["cjk", "東京 李明"],
        ["emoji", "👩🏽‍💻 🧪"],
        ["embedded-bom", "prefix\ufeffsuffix"],
    ]
    document = _document(rows)
    if use_shared_strings:
        document = _with_shared_strings(document)
    document = _transcode_xml_parts(document, xml_encoding)

    [block] = list(XlsxParser().iter_batches(document))

    assert block.metadata["table_row_numbers"] == list(range(2, 8))
    for _, value in rows[1:]:
        assert value in block.text


def test_preserves_error_cells_and_reports_missing_formula_caches():
    doc = _document([["error", "formula"], ["#DIV/0!", "=1+1"]])
    [block] = list(XlsxParser().iter_batches(doc))
    assert "#DIV/0!" in block.text
    assert "[formula: =1+1; cached value unavailable]" in block.text
    assert block.metadata["sheet_formula_cells"] == 1
    assert block.metadata["sheet_missing_formula_values"] == 1
    assert block.metadata["spreadsheet_backend"] == "openpyxl"


def test_cached_formula_values_are_returned_without_execution():
    doc = _document([["formula"], ["=1+1"]])
    doc = _edit_xml(
        doc, lambda xml: xml.replace("<v />", "<v>2</v>").replace("<v/>", "<v>2</v>").replace("<v></v>", "<v>2</v>")
    )
    [block] = list(XlsxParser().iter_batches(doc))
    assert "| 2 [formula: =1+1] |" in block.text
    assert block.metadata["formula_policy"] == "compact_groups_and_individual_formulas"
    assert block.metadata["sheet_formula_cells"] == 1
    assert block.metadata["sheet_missing_formula_values"] == 0


def test_formulas_only_in_hidden_rows_or_columns_do_not_open_formula_reader(monkeypatch):
    import openpyxl

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "value", "hidden helper"])
    sheet.append(["visible", 1, "=B2*2"])
    sheet.append(["hidden", "=1+1", None])
    sheet.column_dimensions["C"].hidden = True
    sheet.row_dimensions[3].hidden = True
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    original_load_workbook = openpyxl.load_workbook
    data_only_calls = []

    def tracking_load_workbook(*args, **kwargs):
        data_only_calls.append(kwargs.get("data_only"))
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", tracking_load_workbook)

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert data_only_calls == [True]
    assert block.metadata["sheet_formula_cells"] == 0
    assert block.metadata["sheet_missing_formula_values"] == 0
    assert "formula:" not in block.text


def test_visible_formula_opens_cached_value_and_formula_readers(monkeypatch):
    import openpyxl

    original_load_workbook = openpyxl.load_workbook
    data_only_calls = []

    def tracking_load_workbook(*args, **kwargs):
        data_only_calls.append(kwargs.get("data_only"))
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", tracking_load_workbook)

    [block] = list(XlsxParser().iter_batches(_document([["formula"], ["=1+1"]])))

    assert data_only_calls == [True, False]
    assert block.metadata["sheet_formula_cells"] == 1
    assert "[formula: =1+1; cached value unavailable]" in block.text


@pytest.mark.parametrize(
    ("formula", "pattern", "semantic", "sources", "unresolved"),
    [
        (
            "=$A2+B$1+$C$3",
            "=RC1+R1C[-1]+R3C3",
            "total = quantity+unit_price+total",
            [
                {"column": "A", "header": "quantity"},
                {"column": "B", "header": "unit_price"},
                {"column": "C", "header": "total"},
            ],
            None,
        ),
        (
            "='Rates Sheet'!$B$2*A2",
            "='Rates Sheet'!R2C2*RC[-2]",
            "total = 'Rates Sheet'!$B$2*quantity",
            [{"column": "A", "header": "quantity"}],
            ["'Rates Sheet'!$B$2"],
        ),
        (
            "=Table1[Amount]*A2",
            "=Table1[Amount]*RC[-2]",
            "total = Table1[Amount]*quantity",
            [{"column": "A", "header": "quantity"}],
            ["Table1[Amount]"],
        ),
    ],
)
def test_formula_analysis_normalizes_references_without_guessing_unresolved_inputs(
    formula, pattern, semantic, sources, unresolved
):
    analysis = module._analyze_formula(
        formula,
        target_row=2,
        target_column=3,
        sheet_name="Data",
        visible_columns=[1, 2, 3],
        headers=["quantity", "unit_price", "total"],
    )

    assert analysis["pattern"] == pattern
    assert analysis["semantic_expression"] == semantic
    assert analysis["source_columns"] == sources
    assert analysis.get("unresolved_references") == unresolved


def test_formula_analysis_bounds_dependency_metadata_for_broad_ranges():
    headers = [f"field_{index}" for index in range(1, 72)]

    analysis = module._analyze_formula(
        "=SUM(A2:BR2)",
        target_row=2,
        target_column=71,
        sheet_name="Data",
        visible_columns=list(range(1, 72)),
        headers=headers,
    )

    assert analysis["pattern"] == "=SUM(RC[-70]:RC[-1])"
    assert analysis["semantic_expression"] == "field_71 = SUM(field_1:field_70)"
    assert len(analysis["source_columns"]) == 64
    assert analysis["source_column_count"] == 70
    assert analysis["source_columns_truncated"] is True


def test_formula_only_rows_remain_searchable_without_cached_results():
    doc = _document([["calculation"], ["=SUM(1,2,3)"]])

    [block] = list(XlsxParser().iter_batches(doc))

    assert "[formula: =SUM(1,2,3); cached value unavailable]" in block.text
    assert block.metadata["table_row_numbers"] == [2]


def test_formula_heavy_sheet_streams_cached_results_and_expressions_across_batches():
    from xml.etree import ElementTree as ET

    rows = [["input", "calculation"]]
    for number in range(1, 251):
        excel_row = number + 1
        rows.append([number, f"=A{excel_row}^2/7"])
    doc = _document(rows)

    def add_cached_results(xml):
        root = ET.fromstring(xml)
        for cell in root.findall(".//{*}c"):
            formula = cell.find("{*}f")
            if formula is None:
                continue
            physical_row = int(cell.get("r")[1:])
            input_value = physical_row - 1
            cell.find("{*}v").text = repr(input_value**2 / 7)
        return ET.tostring(root, encoding="unicode")

    doc = _edit_xml(doc, add_cached_results)
    blocks = list(XlsxParser(batch_size=37).iter_batches(doc))

    assert len(blocks) == 7
    assert [row for block in blocks for row in block.metadata["table_row_numbers"]] == list(range(2, 252))
    assert "| 1 | 0.14285714285714285 |" in blocks[0].text
    assert "| 250 | 8928.57142857143 |" in blocks[-1].text
    assert "[formula:" not in blocks[0].text
    assert blocks[0].metadata["formula_groups"] == [
        {
            "target_column": "B",
            "target_header": "calculation",
            "range": "B2:B38",
            "pattern": "=RC[-1]^2/7",
            "semantic_expression": "calculation = input^2/7",
            "source_columns": [{"column": "A", "header": "input"}],
            "start_row": 2,
            "end_row": 38,
            "formula_count": 37,
        }
    ]
    assert "calculation = input^2/7" in blocks[0].text
    assert all(block.metadata["sheet_formula_cells"] == 250 for block in blocks)
    assert all(block.metadata["sheet_missing_formula_values"] == 0 for block in blocks)


def test_repeated_formulas_are_compacted_with_semantic_dependencies_and_exceptions():
    from xml.etree import ElementTree as ET

    from core.chunking.markdown_utils import MDElement, chunk_table

    rows = [["quantity", "unit_price", "total"]]
    for number in range(1, 26):
        excel_row = number + 1
        suffix = "+1" if excel_row == 13 else ""
        rows.append([number, number * 10, f"=A{excel_row}+B{excel_row}{suffix}"])
    doc = _document(rows)

    def add_cached_results(xml):
        root = ET.fromstring(xml)
        for cell in root.findall(".//{*}c"):
            formula = cell.find("{*}f")
            if formula is None:
                continue
            row_number = int(cell.get("r")[1:])
            number = row_number - 1
            cell.find("{*}v").text = str(number + number * 10 + (1 if row_number == 13 else 0))
        return ET.tostring(root, encoding="unicode")

    [block] = list(XlsxParser().iter_batches(_edit_xml(doc, add_cached_results)))

    assert block.metadata["formula_groups"] == [
        {
            "target_column": "C",
            "target_header": "total",
            "range": "C2:C12",
            "pattern": "=RC[-2]+RC[-1]",
            "semantic_expression": "total = quantity+unit_price",
            "source_columns": [
                {"column": "A", "header": "quantity"},
                {"column": "B", "header": "unit_price"},
            ],
            "start_row": 2,
            "end_row": 12,
            "formula_count": 11,
        },
        {
            "target_column": "C",
            "target_header": "total",
            "range": "C14:C26",
            "pattern": "=RC[-2]+RC[-1]",
            "semantic_expression": "total = quantity+unit_price",
            "source_columns": [
                {"column": "A", "header": "quantity"},
                {"column": "B", "header": "unit_price"},
            ],
            "start_row": 14,
            "end_row": 26,
            "formula_count": 13,
        },
    ]
    [exception] = block.metadata["formula_cells"]
    assert exception["target"] == "C13"
    assert exception["formula"] == "=A13+B13+1"
    assert exception["pattern"] == "=RC[-2]+RC[-1]+1"
    assert exception["semantic_expression"] == "total = quantity+unit_price+1"
    assert "133 [formula: =A13+B13+1]" in block.text
    assert "[formula: =A2+B2]" not in block.text
    assert "C (total), C2:C12: total = quantity+unit_price" in block.text

    chunks = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata),
        chunk_size=90,
        length_function=lambda text: len(text.split()),
    )
    assert len(chunks) > 1
    for chunk in chunks:
        chunk_rows = set(chunk.metadata["table_row_numbers"])
        for group in chunk.metadata.get("formula_groups", []):
            assert any(group["start_row"] <= row <= group["end_row"] for row in chunk_rows)
            assert group["target_column"] == "C"
        for cell in chunk.metadata.get("formula_cells", []):
            assert cell["row"] in chunk_rows


def test_formula_groups_remain_scoped_to_detected_tables_and_worksheets():
    workbook = Workbook()
    first = workbook.active
    first.title = "Operations"
    second = workbook.create_sheet("Finance")
    for sheet in (first, second):
        sheet.append(["quantity", "price", "total"])
        for row_number in range(2, 22):
            sheet.append([row_number, row_number * 10, f"=A{row_number}+B{row_number}"])
    first["E1"], first["F1"], first["G1"] = "hours", "rate", "cost"
    for row_number in range(2, 22):
        first.cell(row_number, 5, row_number)
        first.cell(row_number, 6, row_number * 100)
        first.cell(row_number, 7, f"=E{row_number}*F{row_number}")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert [(block.metadata["sheet_name"], block.metadata["table_id"]) for block in blocks] == [
        ("Operations", "sheet:0:table:0"),
        ("Operations", "sheet:0:table:1"),
        ("Finance", "sheet:1:table:0"),
    ]
    assert [block.metadata["formula_groups"][0]["range"] for block in blocks] == [
        "C2:C21",
        "G2:G21",
        "C2:C21",
    ]
    assert [block.metadata["formula_groups"][0]["semantic_expression"] for block in blocks] == [
        "total = quantity+price",
        "cost = hours*rate",
        "total = quantity+price",
    ]
    assert "Worksheet: Operations" in blocks[0].text
    assert "Worksheet: Operations" in blocks[1].text
    assert "Worksheet: Finance" in blocks[2].text


def test_preserves_precise_numbers_and_mathematical_symbols():
    precise = 0.123456789012345
    notation = "σ² = ∑(xᵢ − μ)² / n; x ≤ μ ± √2; Δ ≈ 10⁻⁶"
    doc = _document([["ratio", "notation"], [precise, notation]])

    [block] = list(XlsxParser().iter_batches(doc))

    assert str(precise) in block.text
    assert notation in block.text


@pytest.mark.parametrize("dimension", ["A1:A1", "A1:XFD1048576"])
def test_untrusted_dimension_neither_truncates_nor_expands_rows(dimension):
    import re

    doc = _document([["name", "city"], ["Alice", "Paris"], ["Bob", "Lyon"]])
    doc = _edit_xml(doc, lambda xml: re.sub(r'<dimension ref="[^"]+"', f'<dimension ref="{dimension}"', xml))
    [block] = list(XlsxParser().iter_batches(doc))
    assert block.metadata["table_row_numbers"] == [2, 3]
    assert block.metadata["table_columns"] == ["name", "city"]
    assert "Bob" in block.text


def test_sparse_rectangle_rejected_before_reader_allocation(monkeypatch):
    doc = _document([["id"], ["value"]])
    doc = _edit_xml(doc, lambda xml: xml.replace('r="2"', 'r="1048576"').replace('r="A2"', 'r="XFD1048576"'))
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))
    with pytest.raises(PipelineError, match="row/column limits"):
        list(XlsxParser().iter_batches(doc))


def test_batch_bytes_bound_accumulation_without_losing_rows():
    doc = _document([["text"], *[["café & < > | " * 10] for _ in range(7)]])
    blocks = list(XlsxParser(batch_size=10_000, limits=XlsxLimits(max_batch_bytes=800)).iter_batches(doc))
    assert len(blocks) > 1
    assert all(len(block.text.encode()) <= 800 for block in blocks)
    assert [row for block in blocks for row in block.metadata["table_row_numbers"]] == list(range(2, 9))


def test_single_row_larger_than_byte_budget_fails_explicitly():
    doc = _document([["text"], ["x" * 200]])
    with pytest.raises(PipelineError, match="batch byte limit"):
        list(XlsxParser(limits=XlsxLimits(max_batch_bytes=100)).iter_batches(doc))


def test_oversized_row_is_rejected_before_markdown_allocation(monkeypatch):
    original = module.markdown_row

    def bounded_markdown_row(cells):
        if any(len(cell) >= 200 for cell in cells):
            pytest.fail("oversized data row reached Markdown rendering")
        return original(cells)

    monkeypatch.setattr(module, "markdown_row", bounded_markdown_row)
    with pytest.raises(PipelineError, match="batch byte limit"):
        list(XlsxParser(limits={"max_batch_bytes": 100}).iter_batches(_document([["text"], ["x" * 200]])))


def test_header_names_are_unique_even_with_generated_name_collisions():
    doc = _document([["column_3", "name", None, "name", "name_2"], [1, 2, 3, 4, 5]])
    [block] = list(XlsxParser().iter_batches(doc))
    assert block.metadata["table_columns"] == ["column_3", "name", "column_3_2", "name_3", "name_2"]


@pytest.mark.parametrize(
    "headers,expected",
    [
        (["Name", "Name "], ["Name", "Name_2"]),
        (["Full\r\nName", "Full\nName"], ["Full\r\nName", "Full\nName_2"]),
    ],
)
def test_rendered_header_names_remain_unique_through_recursive_chunking(headers, expected):
    from core.chunking.recursive import RecursiveSplitter
    from core.models.document import ProcessedDocument

    doc = _document([headers, ["Alice", "Bob"]])
    [block] = list(XlsxParser().iter_batches(doc))
    processed = ProcessedDocument(document_id=doc.id, text_blocks=[block], metadata={"source": doc.filename})

    chunks = RecursiveSplitter(chunk_size=100, length_function=lambda text: len(text.split())).chunk(
        processed, partition="p1"
    )

    assert block.metadata["table_columns"] == expected
    assert chunks
    assert all(chunk.metadata["table_columns"] == expected for chunk in chunks)


def test_whitespace_only_cells_do_not_become_headers_or_data():
    doc = _document([["title"], ["name", "   "], ["Alice", "\t"], [" ", "\t"], ["Bob", "Lyon"]])

    [block] = list(XlsxParser(header_row=2).iter_batches(doc))

    assert block.metadata["table_columns"] == ["name", "column_2"]
    assert block.metadata["table_row_numbers"] == [3, 5]
    assert "Alice" in block.text and "Bob" in block.text


def test_explicit_whitespace_only_header_is_rejected():
    with pytest.raises(PipelineError, match="selected header row 1 is empty"):
        list(XlsxParser(header_row=1).iter_batches(_document([["   ", "\t"], [1, 2]])))


def test_rejects_case_insensitive_duplicate_sheet_names_before_reader_opens(monkeypatch):
    doc = _edit_xml(
        Document(filename="people.xlsx", content_type=DocumentType.XLSX, raw_bytes=_workbook_bytes()),
        lambda xml: xml.replace('name="Projects"', 'name="people"'),
        target="xl/workbook.xml",
    )
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))

    with pytest.raises(PipelineError, match="sheet names are missing or duplicated"):
        list(XlsxParser().iter_batches(doc))


def test_rejects_duplicate_workbook_relationship_ids_before_reader_opens(monkeypatch):
    from xml.etree import ElementTree as ET

    def duplicate_identifier(xml):
        root = ET.fromstring(xml)
        relationship = next(iter(root))
        ET.SubElement(root, relationship.tag, relationship.attrib)
        return ET.tostring(root, encoding="unicode")

    doc = _edit_xml(
        _document([["name"], ["Alice"]]),
        duplicate_identifier,
        target="xl/_rels/workbook.xml.rels",
    )
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))

    with pytest.raises(PipelineError, match="duplicate identifiers"):
        list(XlsxParser().iter_batches(doc))


def test_rejects_two_sheets_pointing_to_the_same_worksheet(monkeypatch):
    from xml.etree import ElementTree as ET

    def duplicate_target(xml):
        root = ET.fromstring(xml)
        worksheets = [node for node in root if node.get("Type", "").endswith("/worksheet")]
        worksheets[1].set("Target", worksheets[0].get("Target"))
        return ET.tostring(root, encoding="unicode")

    doc = _edit_xml(
        Document(filename="people.xlsx", content_type=DocumentType.XLSX, raw_bytes=_workbook_bytes()),
        duplicate_target,
        target="xl/_rels/workbook.xml.rels",
    )
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))

    with pytest.raises(PipelineError, match="duplicate target"):
        list(XlsxParser().iter_batches(doc))


def test_early_consumer_close_closes_workbook(monkeypatch):
    import openpyxl

    original = openpyxl.load_workbook
    closed = []

    def load(*args, **kwargs):
        workbook = original(*args, **kwargs)
        close = workbook.close

        def record_close():
            close()
            closed.append(True)

        workbook.close = record_close
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    batches = XlsxParser(batch_size=1).iter_batches(_document([["a"], [1], [2]]))
    next(batches)
    batches.close()
    assert closed == [True]


@pytest.mark.parametrize(
    "edit",
    [
        lambda xml: xml.replace('r="2"', 'r="1000000000"'),
        lambda xml: xml.replace('r="A2"', 'r="A1"'),
        lambda xml: xml[:-10],
    ],
)
def test_malformed_coordinates_or_xml_raise_domain_error(edit):
    doc = _edit_xml(_document([["a"], [1]]), edit)
    with pytest.raises(PipelineError):
        list(XlsxParser().iter_batches(doc))


def test_mismatched_xml_encoding_is_rejected_before_reader_opens(monkeypatch):
    doc = _edit_xml(
        _document([["name"], ["René"]]),
        lambda xml: '<?xml version="1.0" encoding="UTF-16"?>' + xml,
    )
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))

    with pytest.raises(PipelineError, match="valid XLSX") as error:
        list(XlsxParser().iter_batches(doc))

    assert error.value.code == "INVALID_XLSX"


def test_header_and_sheet_limits():
    doc = _document([["long_header"], [1]])
    with pytest.raises(PipelineError, match="character limit"):
        list(XlsxParser(limits=XlsxLimits(max_cell_chars=3)).iter_batches(doc))
    doc = Document(filename="many.xlsx", content_type=DocumentType.XLSX, raw_bytes=_workbook_bytes())
    with pytest.raises(PipelineError, match="worksheet count"):
        list(XlsxParser(limits=XlsxLimits(max_sheets=1)).iter_batches(doc))


@pytest.mark.parametrize(
    "headers,values",
    [
        (["question", "answer"], ["question " * 780, "answer " * 2000]),
        (["id", "name"], ["identifier " * 2000, "person " * 4000]),
        (["a|b\nc", "other"], ["value " * 2000, "other " * 2000]),
    ],
)
def test_oversized_xlsx_cells_are_lossless_bounded_and_have_exact_citations(headers, values):
    from core.chunking.markdown_utils import MDElement, chunk_table
    from core.indexing.parsers.tabular.base import markdown_row

    [block] = list(XlsxParser().iter_batches(_document([headers, values])))
    chunks = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata),
        chunk_size=300,
        length_function=lambda text: len(text.split()),
    )
    assert 1 < len(chunks) < 100
    assert all(len(chunk.content.split()) <= 300 for chunk in chunks)
    for column_index, value in enumerate(values, start=1):
        parts = []
        for chunk in chunks:
            if chunk.metadata["table_column_indices"] != [column_index]:
                continue
            assert chunk.metadata["table_row_numbers"] == [2]
            letter = chr(64 + column_index)
            assert chunk.metadata["cell_range"] == f"'Sheet'!{letter}2:{letter}2"
            body = chunk.content.splitlines()[2].split("|")[1].strip()
            if " continuation " in body:
                body = body.split("] ", 1)[1]
            parts.append(body)
        expected = markdown_row([value]).split("|")[1].strip()
        # Markdown rendering trims cell-edge whitespace; internal words survive.
        assert " ".join(parts).split() == expected.split()


def test_xlsx_multirow_chunk_citations_do_not_retain_batch_end():
    from core.chunking.markdown_utils import MDElement, chunk_table

    [block] = list(XlsxParser().iter_batches(_document([["name"], ["Alice"], ["Bob"], ["Claire"]])))
    chunks = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata), chunk_size=50, length_function=len
    )
    assert len(chunks) == 3
    assert [chunk.metadata["cell_range"] for chunk in chunks] == [f"'Sheet'!A{i}:A{i}" for i in (2, 3, 4)]


def _with_shared_strings(document):
    import zipfile
    from xml.etree import ElementTree as ET

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    strings = []
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(document.raw_bytes)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            content = source.read(name)
            root = ET.fromstring(content) if name.endswith((".xml", ".rels")) else None
            if name == "xl/worksheets/sheet1.xml":
                for cell in root.findall(".//{*}c"):
                    if cell.get("t") != "inlineStr":
                        continue
                    text = "".join(t.text or "" for t in cell.findall(".//{*}t"))
                    if text not in strings:
                        strings.append(text)
                    cell.remove(cell.find("{*}is"))
                    cell.set("t", "s")
                    ET.SubElement(cell, f"{{{namespace}}}v").text = str(strings.index(text))
            elif name == "[Content_Types].xml":
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName="/xl/sharedStrings.xml",
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                )
            elif name == "xl/_rels/workbook.xml.rels":
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="rIdShared",
                    Target="sharedStrings.xml",
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                )
            if root is not None:
                content = ET.tostring(root)
            dest.writestr(name, content)
        root = ET.Element(f"{{{namespace}}}sst")
        for text in strings:
            ET.SubElement(ET.SubElement(root, f"{{{namespace}}}si"), f"{{{namespace}}}t").text = text
        dest.writestr("xl/sharedStrings.xml", ET.tostring(root))
    return document.model_copy(update={"raw_bytes": output.getvalue()})


def test_shared_string_count_is_bounded_before_openpyxl_loads_sheet(monkeypatch):
    doc = _with_shared_strings(_document([["text"], *[["long shared text " * 10] for _ in range(10)]]))
    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))
    with pytest.raises(PipelineError, match="shared string count"):
        list(XlsxParser(limits=XlsxLimits(max_shared_strings=1)).iter_batches(doc))


def test_shared_strings_are_preserved():
    doc = _with_shared_strings(_document([["name", "city"], ["René", "Paris"], ["René", "Paris"]]))
    [block] = list(XlsxParser().iter_batches(doc))
    assert block.text.count("René") == 2


@pytest.mark.parametrize("strategy", ["recursive", "structured"])
def test_both_chunkers_preserve_xlsx_identity_across_sheets_and_split_cells(strategy):
    from core.chunking.recursive import RecursiveSplitter
    from core.chunking.structured_section import StructuredSectionChunker
    from core.models.chunk import ChunkType
    from core.models.document import ProcessedDocument

    workbook = Workbook()
    workbook.active.title = "First"
    for sheet in (workbook.active, workbook.create_sheet("Second")):
        sheet.append(["a|b\nc", "name"])
        sheet.append([f"{sheet.title} " * 1500, "person " * 2000])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    doc = Document(id="same-file", filename="source.xlsx", content_type=DocumentType.XLSX, raw_bytes=output.getvalue())
    blocks = list(XlsxParser().iter_batches(doc))
    processed = ProcessedDocument(document_id=doc.id, text_blocks=blocks, metadata={"source": doc.filename})

    def count(text):
        return len(text.split())

    chunker = (
        RecursiveSplitter(chunk_size=300, length_function=count)
        if strategy == "recursive"
        else StructuredSectionChunker(chunk_size=200, max_tokens=300, min_tokens=0, length_function=count)
    )
    chunks = chunker.chunk(processed, partition="p1")
    assert 4 < len(chunks) < 100
    assert {c.metadata["table_id"] for c in chunks} == {"sheet:0:table:0", "sheet:1:table:0"}
    for chunk in chunks:
        assert chunk.chunk_type == ChunkType.TABLE
        assert chunk.document_id == doc.id and chunk.partition == "p1"
        assert chunk.token_count <= 300
        assert chunk.metadata["table_row_numbers"] == [2]
        assert chunk.metadata["cell_range"].endswith(("A2:A2", "B2:B2"))
        assert chunk.metadata["cell_range"].startswith(f"'{chunk.metadata['sheet_name']}'!")
        assert f"Worksheet: {chunk.metadata['sheet_name']}" in chunk.text
        assert "XLSX row 2" in chunk.text
        assert "CSV row" not in chunk.text
        assert "&#124;" in chunk.text or chunk.metadata["table_columns"] == ["name"]
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


def test_explicit_header_skips_report_title_and_preserves_physical_citations():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["Annual report"])
    sheet.merge_cells("A1:B1")
    sheet.append([])
    sheet.append(["name", "city"])
    sheet.append(["Alice", "Paris"])
    sheet.append([])
    sheet.append(["René", "Lyon"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    blocks = list(XlsxParser(header_row=3, batch_size=1).iter_batches(Document(raw_bytes=output.getvalue())))
    assert [b.metadata["table_row_numbers"] for b in blocks] == [[4], [6]]
    assert all(b.metadata["table_columns"] == ["name", "city"] for b in blocks)
    assert all(b.metadata["table_header_range"] == "'Report'!A3:B3" for b in blocks)
    assert blocks[1].metadata["cell_range"] == "'Report'!A6:B6"
    assert all("Annual report" in b.text for b in blocks)
    assert all(b.metadata["sheet_context_row_numbers"] == [1] for b in blocks)
    assert "René" in blocks[1].text


@pytest.mark.parametrize("header_row, message", [(1, "empty"), (99, "missing")])
def test_invalid_selected_header_is_not_silently_inferred(header_row, message):
    with pytest.raises(PipelineError, match=message):
        list(XlsxParser(header_row=header_row).iter_batches(Document(raw_bytes=_workbook_bytes())))


@pytest.mark.parametrize("header_row", [0, -1, True, 1.5, "3", 1_048_577])
def test_header_row_requires_a_valid_physical_excel_row(header_row):
    from core.utils.exceptions import ConfigError

    with pytest.raises(ConfigError):
        XlsxParser(header_row=header_row)


def test_sheet_headers_override_global_and_allow_automatic_selection():
    workbook = Workbook()
    workbook.active.title = "People"
    for sheet in (workbook.active, workbook.create_sheet("Projects"), workbook.create_sheet("Notes")):
        sheet.append(["title"])
        sheet.append(["name", "value"])
        sheet.append(["Alice", "Paris"])
        sheet.append(["René", "Lyon"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    mapping = {"People": 3, "Notes": None}
    parser = XlsxParser(batch_size=1, header_row=2, sheet_header_rows=mapping)
    mapping["People"] = 4  # The parser owns a snapshot of its configuration.
    blocks = list(parser.iter_batches(Document(raw_bytes=output.getvalue())))
    by_sheet = {
        name: [b for b in blocks if b.metadata["sheet_name"] == name] for name in ["People", "Projects", "Notes"]
    }
    assert by_sheet["People"][0].metadata["table_columns"] == ["Alice", "Paris"]
    assert by_sheet["People"][0].metadata["table_header_range"] == "'People'!A3:B3"
    assert by_sheet["People"][0].metadata["cell_range"] == "'People'!A4:B4"
    assert [b.metadata["table_row_numbers"] for b in by_sheet["Projects"]] == [[3], [4]]
    # Without a blank structural boundary, automatic detection keeps the
    # first nonempty row rather than risking loss of an ordinary data record.
    assert by_sheet["Notes"][0].metadata["table_header_row"] == 1


@pytest.mark.parametrize("overrides", [{"people": 2}, {"Hidden": 1}, {"People": 99}])
def test_sheet_override_validation_precedes_backend_and_emission(overrides, monkeypatch):
    parser = XlsxParser(sheet_header_rows=overrides)

    def unexpected(*args):
        pytest.fail("invalid overrides must fail before opening a reader")

    monkeypatch.setattr(module, "_openpyxl_sheets", unexpected)
    with pytest.raises(PipelineError):
        next(parser.iter_batches(Document(raw_bytes=_workbook_bytes())))


def test_sheet_override_empty_header_is_rejected():
    with pytest.raises(PipelineError, match="empty"):
        list(XlsxParser(sheet_header_rows={"People": 1}).iter_batches(Document(raw_bytes=_workbook_bytes())))


@pytest.mark.parametrize(
    "overrides",
    [
        [],
        {"": 1},
        {"  ": 1},
        {1: 2},
        {"People": True},
        {"People": 0},
        {"People": "2"},
        {"People": 1.5},
        {"People": 1_048_577},
    ],
)
def test_invalid_sheet_override_configuration(overrides):
    from core.utils.exceptions import ConfigError

    with pytest.raises(ConfigError):
        XlsxParser(sheet_header_rows=overrides)


def test_loader_config_validates_sheet_rows():
    from core.config.indexation import LoaderConfig
    from pydantic import ValidationError

    config = LoaderConfig(xlsx_sheet_header_rows={"People": 3, "Notes": None})
    assert config.xlsx_sheet_header_rows == {"People": 3, "Notes": None}
    for invalid_row in (True, 0, "3", 1_048_577):
        with pytest.raises(ValidationError):
            LoaderConfig(xlsx_sheet_header_rows={"People": invalid_row})


@pytest.mark.parametrize(
    "schemas",
    [
        [],
        {"": ["id"]},
        {"Raw": []},
        {"Raw": [""]},
        {"Raw": [1]},
        {1: ["id"]},
    ],
)
def test_invalid_headerless_schema_configuration(schemas):
    from core.utils.exceptions import ConfigError

    with pytest.raises(ConfigError):
        XlsxParser(sheet_headerless_schemas=schemas)


def test_headerless_schema_rejects_conflicting_or_wrong_width_configuration():
    from core.utils.exceptions import ConfigError

    with pytest.raises(ConfigError, match="cannot be combined"):
        XlsxParser(header_row=1, sheet_headerless_schemas={"Raw": ["id"]})
    with pytest.raises(ConfigError, match="cannot be combined"):
        XlsxParser(sheet_header_rows={"Raw": 1}, sheet_headerless_schemas={"Raw": ["id"]})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Raw"
    sheet.append(["A", "B"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(PipelineError, match="headerless schema has 1 columns, expected 2"):
        list(
            XlsxParser(sheet_headerless_schemas={"Raw": ["only_one"]}).iter_batches(
                Document(raw_bytes=output.getvalue())
            )
        )


@pytest.mark.parametrize("repeated", [False, True])
def test_auto_header_skips_title_and_preserves_rows(repeated):
    wb = Workbook()
    ws = wb.active
    ws.append(["Quarterly report"] * (3 if repeated else 1))
    ws.append([])
    ws.append(["id", "name", "summary"])
    ws.append(["R-1", "Alice", "First"])
    ws.append(["R-2", "Bob", "Second"])
    output = BytesIO()
    wb.save(output)
    blocks = list(XlsxParser(batch_size=1).iter_batches(Document(raw_bytes=output.getvalue())))
    assert [b.metadata["table_row_numbers"] for b in blocks] == [[4], [5]]
    assert all(b.metadata["table_header_row"] == 3 for b in blocks)
    assert all(b.metadata["table_columns"] == ["id", "name", "summary"] for b in blocks)


def test_auto_header_skips_multiple_narrow_preamble_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Regional report", "Quarter 4"])
    sheet.append(["Prepared by", "Accessibility team"])
    sheet.append([])
    sheet.append(["id", "service", "region", "policy"])
    sheet.append(["R-1", "Screen reader desk", "West", "Contact allowed"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_header_row"] == 4
    assert block.metadata["table_row_numbers"] == [5]
    assert block.metadata["sheet_context_row_numbers"] == [1, 2]
    assert block.metadata["table_columns"] == ["id", "service", "region", "policy"]


def test_auto_header_preserves_single_column_and_explicit_title():
    wb = Workbook()
    ws = wb.active
    ws.append(["name"])
    ws.append(["Alice"])
    ws.append(["Bob"])
    output = BytesIO()
    wb.save(output)
    block = next(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))
    assert block.metadata["table_header_row"] == 1
    assert block.metadata["table_row_numbers"] == [2, 3]


def test_auto_header_scan_is_bounded_and_replays_all_rows():
    rows = iter([["title"]] + [[""]] * 31 + [["id", "name"], ["1", "Alice"]])
    selected, replay = module._detect_header(rows, 1024)
    assert selected == 1
    assert len(list(replay)) == 34
    rows = iter([["title"], ["id", "name"], ["1", "Alice"]])
    selected, replay = module._detect_header(rows, 1)
    assert selected == 1
    assert list(replay) == [["title"], ["id", "name"], ["1", "Alice"]]


def test_auto_header_keeps_ambiguous_header_and_does_not_skip_numeric_data():
    original = [["name", "name"], [1, 2], [3, 4]]
    selected, replay = module._detect_header(iter(original), 1024)
    assert selected == 1
    assert list(replay) == original


def test_auto_header_keeps_equal_width_text_rows_ambiguous():
    original = [["code", "label"], ["A-1", "First"], ["A-2", "Second"]]
    selected, replay = module._detect_header(iter(original), 1024)
    assert selected == 1
    assert list(replay) == original


def test_auto_header_does_not_consume_a_narrow_first_data_record():
    document = _document([["record"], ["HL-001", "Alice"], ["HL-002", "Bob"]])

    [block] = list(XlsxParser().iter_batches(document))

    assert block.metadata["table_header_row"] == 1
    assert block.metadata["table_row_numbers"] == [2, 3]
    assert "HL-001" in block.text


@pytest.mark.parametrize("target", ["xl/theme/theme1.xml", "docProps/core.xml"])
def test_eager_openpyxl_metadata_is_bounded_before_reader_opens(target, monkeypatch):
    document = _replace_part(_document([["id"], [1]]), target, bytes(range(251)) * 100)

    def unexpected_reader(*args):
        pytest.fail("Oversized eager metadata must be rejected before opening the reader")

    monkeypatch.setattr(module, "_openpyxl_sheets", unexpected_reader)
    with pytest.raises(PipelineError, match="metadata exceeds the memory limit") as error:
        list(XlsxParser(limits={"max_metadata_bytes": 20_000}).iter_batches(document))
    assert error.value.code == "XLSX_LIMIT_EXCEEDED"


def test_aggregate_eager_metadata_is_bounded_before_reader_opens(monkeypatch):
    document = _document([["id"], [1]])
    document = _replace_part(document, "xl/theme/theme1.xml", bytes(range(251)) * 60)
    document = _replace_part(document, "docProps/core.xml", bytes(range(250, -1, -1)) * 60)

    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))
    with pytest.raises(PipelineError, match="metadata exceeds the memory limit") as error:
        list(XlsxParser(limits={"max_metadata_bytes": 30_000}).iter_batches(document))
    assert error.value.code == "XLSX_LIMIT_EXCEEDED"


def test_worksheet_relationship_metadata_is_bounded_before_reader_opens(monkeypatch):
    import zipfile

    document = _document([["id"], [1]])
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(document.raw_bytes)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for name in source.namelist():
            dest.writestr(name, source.read(name))
        dest.writestr("xl/worksheets/_rels/sheet1.xml.rels", bytes(range(251)) * 120)
    document = document.model_copy(update={"raw_bytes": output.getvalue()})

    monkeypatch.setattr(module, "_openpyxl_sheets", lambda *_: pytest.fail("reader must not open"))
    with pytest.raises(PipelineError, match="metadata exceeds the memory limit") as error:
        list(XlsxParser(limits={"max_metadata_bytes": 30_000}).iter_batches(document))
    assert error.value.code == "XLSX_LIMIT_EXCEEDED"


@pytest.mark.parametrize("target", ["xl/styles.xml", "docProps/core.xml"])
def test_malformed_backend_metadata_raises_domain_error(target):
    document = _replace_part(_document([["id"], [1]]), target, b"<broken")

    with pytest.raises(PipelineError, match="valid XLSX workbook") as error:
        list(XlsxParser().iter_batches(document))
    assert error.value.code == "INVALID_XLSX"
    assert error.value.status_code == 422


@pytest.mark.parametrize("strategy", ["recursive", "structured"])
@pytest.mark.parametrize("header_row", [None, 3])
def test_preamble_is_searchable_in_each_split_chunk(strategy, header_row):
    from core.chunking.recursive import RecursiveSplitter
    from core.chunking.structured_section import StructuredSectionChunker
    from core.models.document import ProcessedDocument

    wb = Workbook()
    ws = wb.active
    ws.append(["Quarterly | report", "Quarterly | report"])
    ws.append([])
    ws.append(["id", "description"])
    for i in range(8):
        ws.append([str(i), "Long description " * 100])
    output = BytesIO()
    wb.save(output)
    blocks = list(XlsxParser(header_row=header_row, batch_size=2).iter_batches(Document(raw_bytes=output.getvalue())))
    assert all(b.metadata["sheet_context_row_numbers"] == [1] for b in blocks)
    assert all(b.text.count("Quarterly | report") == 1 for b in blocks)
    processed = ProcessedDocument(document_id="test", text_blocks=blocks, metadata={"source": "report.xlsx"})

    def count(text):
        return len(text.split())

    chunker = (
        RecursiveSplitter(chunk_size=100, length_function=count)
        if strategy == "recursive"
        else StructuredSectionChunker(chunk_size=80, max_tokens=100, min_tokens=0, length_function=count)
    )
    chunks = chunker.chunk(processed, partition="default")
    assert len(chunks) > 4
    for chunk in chunks:
        assert "Quarterly | report" in chunk.text
        assert chunk.token_count <= 100
        assert chunk.metadata["sheet_context_row_numbers"] == [1]
        assert all(r >= 4 for r in chunk.metadata["table_row_numbers"])


def test_blank_preamble_does_not_create_context():
    blocks = list(XlsxParser(header_row=2).iter_batches(Document(raw_bytes=_workbook_bytes())))
    assert "sheet_context" not in blocks[0].metadata
    assert blocks[0].metadata["worksheet_context"] == "\n\nWorksheet: People"
    assert blocks[0].text.endswith("\n\nWorksheet: People")


def test_explicit_preamble_retains_multiple_rows_and_values():
    wb = Workbook()
    ws = wb.active
    ws.append(["Report", "Region: West"])
    ws.append(["Period: 2026", "Period: 2026"])
    ws.append(["id", "value"])
    ws.append([1, 2])
    output = BytesIO()
    wb.save(output)
    block = next(XlsxParser(header_row=3).iter_batches(Document(raw_bytes=output.getvalue())))
    assert block.metadata["sheet_context_row_numbers"] == [1, 2]
    assert "Report ; Region: West" in block.text
    assert block.text.count("Period: 2026") == 1


def test_excessive_preamble_fails_without_truncation():
    wb = Workbook()
    ws = wb.active
    ws.append(["Title " * 100])
    ws.append(["id", "value"])
    ws.append([1, 2])
    output = BytesIO()
    wb.save(output)
    with pytest.raises(PipelineError, match="context exceeds"):
        list(
            XlsxParser(header_row=2, limits={"max_batch_bytes": 100}).iter_batches(
                Document(raw_bytes=output.getvalue())
            )
        )


def test_sheet_context_larger_than_chunk_budget_fails():
    from core.chunking.markdown_utils import MDElement, chunk_table

    context = "\n\nWorksheet context:\n" + "title " * 100
    table = "| id |\n| --- |\n| 1 |"
    element = MDElement(
        type="table",
        content=table + context,
        metadata={"tabular_format": "xlsx", "sheet_context": context, "table_columns": ["id"], "csv_columns": ["id"]},
    )
    with pytest.raises(PipelineError, match="context exceeds chunk token budget"):
        chunk_table(element, chunk_size=20, length_function=lambda text: len(text.split()))


def test_xlsx_continuation_restores_context_to_expanded_chunk_budget():
    from core.chunking.markdown_utils import MDElement, chunk_table

    column = " ".join(f"header{number}" for number in range(12))
    value = " ".join(f"value{number}" for number in range(4))
    context = "\n\nWorksheet context:\nQuarterly report"
    table = f"| {column} |\n| --- |\n| {value} |"
    element = MDElement(
        type="table",
        content=table + context,
        metadata={
            "tabular_format": "xlsx",
            "sheet_context": context,
            "sheet_name": "Sheet",
            "table_columns": [column],
            "csv_columns": [column],
            "table_row_start": 2,
            "table_row_end": 2,
            "table_row_numbers": [2],
        },
    )

    def count_words(text):
        return len(text.split())

    chunks = chunk_table(element, chunk_size=10, length_function=count_words)

    assert len(chunks) == 4
    assert all(chunk.content.endswith(context) for chunk in chunks)
    assert all(count_words(chunk.content) > 10 for chunk in chunks)
    assert all(
        count_words(chunk.content)
        <= chunk.metadata["csv_effective_chunk_size"] + count_words("\n") + count_words(context)
        for chunk in chunks
    )


@pytest.mark.parametrize("state", ["", "Visible", "veryhidden", "invalid"])
def test_invalid_sheet_visibility_is_rejected_before_reader_opens(state, monkeypatch):
    document = _edit_xml(
        _document([["id", "value"], [1, "kept"]]),
        lambda xml: xml.replace('state="visible"', f'state="{state}"'),
        "xl/workbook.xml",
    )

    def unexpected_reader(*args):
        pytest.fail("Invalid workbook metadata must be rejected before opening the reader")

    monkeypatch.setattr(module, "_openpyxl_sheets", unexpected_reader)
    with pytest.raises(PipelineError, match="invalid visibility state") as error:
        list(XlsxParser().iter_batches(document))
    assert error.value.code == "INVALID_XLSX"


def test_very_hidden_sheet_is_excluded_and_false_zero_rows_keep_exact_citations():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "O'Brien directory"
    sheet.append(["count", "active"])
    sheet.append([0, False])
    sheet.append([])
    sheet.append([1, True])
    hidden = workbook.create_sheet("Private")
    hidden.sheet_state = "veryHidden"
    hidden.append(["secret"])
    hidden.append(["must not be indexed"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = list(XlsxParser(batch_size=1).iter_batches(Document(raw_bytes=output.getvalue())))

    assert len(blocks) == 2
    assert "| 0 | false |" in blocks[0].text
    assert "| 1 | true |" in blocks[1].text
    assert [block.metadata["table_row_numbers"] for block in blocks] == [[2], [4]]
    assert [block.metadata["cell_range"] for block in blocks] == [
        "'O''Brien directory'!A2:B2",
        "'O''Brien directory'!A4:B4",
    ]
    assert all(block.metadata["hidden_sheets_skipped"] == 1 for block in blocks)
    assert all("must not be indexed" not in block.text for block in blocks)


def test_hidden_rows_and_columns_are_not_indexed():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "internal_note", "public_value"])
    sheet.append(["R-1", "row secret", "hidden row value"])
    sheet.append(["R-2", "column secret", "visible value"])
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_columns"] == ["id", "public_value"]
    assert block.metadata["table_row_numbers"] == [3]
    assert "visible value" in block.text
    assert "row secret" not in block.text
    assert "hidden row value" not in block.text
    assert "internal_note" not in block.text
    assert "column secret" not in block.text
    assert block.metadata["hidden_rows_skipped"] == 1
    assert block.metadata["hidden_columns_skipped"] == 1
    assert block.metadata["table_worksheet_column_indices"] == [1, 3]


def test_hidden_columns_keep_physical_cell_coordinates_after_chunking():
    from core.chunking.markdown_utils import MDElement, chunk_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["id", "secret", "description"])
    sheet.append(["R-1", "classified", "public"])
    sheet.column_dimensions["B"].hidden = True
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))
    chunks = chunk_table(
        MDElement(type="table", content=block.text, metadata=block.metadata),
        chunk_size=200,
        length_function=lambda text: len(text.split()),
    )

    assert len(chunks) == 1
    assert chunks[0].metadata["cell_range"] == "'Report'!A2:C2"
    assert chunks[0].metadata["table_worksheet_column_indices"] == [1, 3]
    assert "classified" not in chunks[0].content


def test_hidden_header_override_is_rejected():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "name"])
    sheet.append(["R-1", "Alice"])
    sheet.row_dimensions[1].hidden = True
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(PipelineError, match="selected header row 1 is hidden"):
        list(XlsxParser(header_row=1).iter_batches(Document(raw_bytes=output.getvalue())))


def test_hidden_unused_columns_beyond_parser_limit_do_not_reject_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id"])
    sheet.append(["R-1"])
    sheet.column_dimensions.group("B", "XFD", hidden=True)
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    [block] = list(XlsxParser().iter_batches(Document(raw_bytes=output.getvalue())))

    assert block.metadata["table_columns"] == ["id"]
    assert block.metadata["table_row_numbers"] == [2]
    assert block.metadata["hidden_columns_skipped"] == 4095
