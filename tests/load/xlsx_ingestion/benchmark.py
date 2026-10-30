"""Local, deterministic parser benchmark; never calls an embedding/LLM endpoint.

Run with a Python environment containing the project and XLSX dependencies:
  PYTHONPATH=openrag python tests/load/xlsx_ingestion/benchmark.py --output /tmp/xlsx-audit.json
Fixtures are created in separate processes, outside each parser's RSS measurement.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import platform
import resource
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

PROFILES = {
    "rows-10000": (10000, 4, 40),
    "rows-100000": (100000, 4, 40),
    "shared-50000": (50000, 4, 40),
    "wide": (200, 256, 40),
    "long-cells": (1000, 4, 2048),
}


def generate(path, profile):
    from openpyxl import Workbook

    count, width, length = PROFILES[profile]
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("People")
    sheet.append([f"field_{i}" for i in range(width)])
    for row in range(count):
        sheet.append(
            [
                "".join(
                    hashlib.sha256(f"{row}:{col}:{part}".encode()).hexdigest() for part in range((length + 63) // 64)
                )[:length]
                for col in range(width)
            ]
        )
    workbook.save(path)
    workbook.close()
    if profile.startswith("shared-"):
        convert_shared_strings(path)


def convert_shared_strings(path):
    """Build an Excel-style shared-string fixture in the generator process only."""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    strings = []
    converted = path.with_suffix(".shared.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(converted, "w", zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                for cell in root.findall(".//{*}c"):
                    inline = cell.find("{*}is")
                    if inline is None:
                        continue
                    strings.append("".join(t.text or "" for t in inline.iter() if t.tag.endswith("}t")))
                    cell.remove(inline)
                    cell.set("t", "s")
                    ET.SubElement(cell, f"{{{namespace}}}v").text = str(len(strings) - 1)
                content = ET.tostring(root)
            elif name == "[Content_Types].xml":
                root = ET.fromstring(content)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName="/xl/sharedStrings.xml",
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                )
                content = ET.tostring(root)
            elif name == "xl/_rels/workbook.xml.rels":
                root = ET.fromstring(content)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="rIdShared",
                    Target="sharedStrings.xml",
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                )
                content = ET.tostring(root)
            dest.writestr(name, content)
        with dest.open("xl/sharedStrings.xml", "w") as stream:
            stream.write(f'<sst xmlns="{namespace}">'.encode())
            for text in strings:
                entry = ET.Element("si")
                ET.SubElement(entry, "t").text = text
                stream.write(ET.tostring(entry))
            stream.write(b"</sst>")
    converted.replace(path)


def measure(path):
    from core.indexing.parsers.tabular.xlsx_parser import XlsxParser
    from core.models.document import Document, DocumentType

    parser = XlsxParser()
    document = Document(filename=path.name, source_path=str(path), content_type=DocumentType.XLSX)
    started = time.perf_counter()
    batches = rows = max_bytes = 0
    digest = hashlib.sha256()
    for block in parser.iter_batches(document):
        batches += 1
        rows += len(block.metadata["table_row_numbers"])
        encoded = block.text.encode()
        digest.update(encoded)
        max_bytes = max(max_bytes, len(encoded))
    elapsed = time.perf_counter() - started
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    peak_mib = peak / (1024**2 if sys.platform == "darwin" else 1024)
    return {
        "backend": "openpyxl",
        "seconds": round(elapsed, 3),
        "peak_rss_mib": round(peak_mib, 2),
        "rows": rows,
        "batches": batches,
        "largest_batch_bytes": max_bytes,
        "sha256": digest.hexdigest(),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("/tmp/xlsx-audit.json"))
    parser.add_argument("--repeats", type=int, default=2)
    parser.add_argument("--generate", choices=PROFILES)
    parser.add_argument("--measure", action="store_true")
    parser.add_argument("--file", type=Path)
    args = parser.parse_args()
    if args.generate:
        generate(args.file, args.generate)
        return
    if args.measure:
        print(json.dumps(measure(args.file)))
        return
    results = []
    script = str(Path(__file__).resolve())
    with tempfile.TemporaryDirectory(prefix="xlsx-audit-") as temp:
        for profile in PROFILES:
            path = Path(temp) / f"{profile}.xlsx"
            subprocess.run([sys.executable, script, "--generate", profile, "--file", str(path)], check=True)
            for repeat in range(args.repeats):
                result = subprocess.run(
                    [sys.executable, script, "--measure", "--file", str(path)],
                    capture_output=True,
                    text=True,
                )
                if result.returncode:
                    raise RuntimeError(result.stderr)
                record = json.loads(result.stdout)
                record.update(profile=profile, repeat=repeat + 1, archive_bytes=path.stat().st_size)
                assert record["rows"] == PROFILES[profile][0]
                results.append(record)
                print(f"{profile}: {record['seconds']} s, {record['peak_rss_mib']} MiB", flush=True)
            assert len({r["sha256"] for r in results if r["profile"] == profile}) == 1, "Output mismatch"
    report = {
        "platform": platform.platform(),
        "python": sys.version,
        "architecture": platform.machine(),
        "libraries": {"openpyxl": importlib.metadata.version("openpyxl")},
        "runs": results,
    }
    args.output.write_text(json.dumps(report, indent=2) + "\n")
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
